from pathlib import Path
import shutil
import zipfile

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
OUT_DIR = ROOT / "outputs" / "luad_signature_extraction"
PKG = OUT_DIR / "manuscript_package"


INPUTS = {
    "full_inventory": OUT_DIR / "LUAD_reproducible_signature_inventory.xlsx",
    "selection": WORK / "luad_geo_manuscript_ready_signature_selection.csv",
    "integrated": WORK / "integrated_main_signature_manuscript_evidence.csv",
    "geo_pairwise": WORK / "luad_geo_manuscript_pairwise_evidence.csv",
    "geo_signature": WORK / "luad_geo_manuscript_signature_evidence.csv",
    "robustness_summary": WORK / "luad_geo_main_signature_robustness_summary.csv",
    "robustness_adjusted": WORK / "luad_geo_main_signature_clinical_adjusted_cox.csv",
    "robustness_zscore": WORK / "luad_geo_main_signature_zscore_sensitivity.csv",
    "tcga_results": WORK / "tcga_luad_main_signature_gdc_star_tpm_sanity_results.csv",
    "cohort_source": WORK / "manuscript_geo_cohort_source_table.csv",
    "loco": WORK / "main_signature_leave_one_cohort_out.csv",
    "a097_note": OUT_DIR / "A097_uploaded_pdf_manual_verification.md",
    "gse50081_note": OUT_DIR / "GSE50081_time_unit_verification.md",
    "core_text": OUT_DIR / "LUAD_manuscript_core_text_draft.md",
    "full_draft": OUT_DIR / "LUAD_full_manuscript_working_draft.md",
    "integrated_note": OUT_DIR / "LUAD_integrated_manuscript_recommendation.md",
    "methods_note": OUT_DIR / "Manuscript_data_sources_and_reproducibility.md",
    "heatmap": OUT_DIR / "LUAD_GEO_main_signature_heatmap.svg",
    "forest": OUT_DIR / "LUAD_GEO_main_signature_pooled_HR_forest.svg",
}


def numeric_field(field):
    return any(
        token in field
        for token in [
            "N",
            "n",
            "HR",
            "p",
            "q",
            "C_index",
            "events",
            "cohorts",
            "count",
            "Delta",
            "I2",
            "rate",
            "Cutoff",
            "Gene_",
        ]
    )


def write_xlsx(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            ws.append(list(row))
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(df.columns, start=1):
            width = 17
            if any(x in col for x in ["Title", "Signature", "Evidence", "limitation", "caveat", "Role", "Decision", "GSEs", "genes"]):
                width = 44
            ws.column_dimensions[get_column_letter(idx)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    wb.save(path)


def table1(selection):
    keep_tiers = [
        "Tier_1_strict_reported_cutoff",
        "Tier_2_primary_median_fallback",
        "Tier_3_secondary_median_fallback",
        "Tier_4_continuous_score_support",
        "Tier_5_manual_override_sensitivity",
    ]
    cols = [
        "Article_ID",
        "PMID",
        "Title",
        "Final_analysis_tier",
        "Recommended_manuscript_role",
        "Gene_count",
        "Cutoff_type",
        "Cutoff_value",
        "Training_cohort",
        "Validation_cohort",
        "Tested_cohorts_ok",
        "Group_sig_p05_n",
        "Continuous_sig_p05_n",
        "Median_C_index",
        "Key_caveats_for_manuscript",
    ]
    t = selection[selection["Final_analysis_tier"].isin(keep_tiers)].copy()
    return t[[c for c in cols if c in t.columns]]


def make_workflow_svg(path):
    width, height = 1280, 720
    boxes = [
        (60, 96, 210, 82, "Full-text extraction", "178 LUAD-related articles screened"),
        (330, 96, 210, 82, "Reproducible formulas", "57 included reproducible signatures"),
        (600, 96, 210, 82, "GEO expression + OS", "9 OS-complete GEO cohorts"),
        (870, 96, 250, 82, "First-pass scoring", "Formula score = sum(coef x expression)"),
        (190, 300, 250, 92, "Manuscript triage", "Strict cutoff vs median fallback vs sensitivity"),
        (515, 300, 250, 92, "Primary GEO portability", "10 main candidates, 75 tests"),
        (840, 300, 250, 92, "Robustness analyses", "Clinical adjustment, BH-FDR, gene z-score"),
        (355, 520, 250, 92, "TCGA sanity check", "Formula direction, not independent validation"),
        (690, 520, 300, 92, "Manuscript evidence package", "Main tables, figures, and cautious claims"),
    ]
    arrows = [
        ((270, 137), (330, 137)),
        ((540, 137), (600, 137)),
        ((810, 137), (870, 137)),
        ((995, 178), (315, 300)),
        ((440, 346), (515, 346)),
        ((765, 346), (840, 346)),
        ((965, 392), (480, 520)),
        ((605, 566), (690, 566)),
    ]
    colors = ["#D9EAF7", "#E8F4E5", "#FFF2CC", "#FCE4D6", "#EADCF8", "#DDEBF7", "#E2F0D9", "#F8CBAD", "#D9EAD3"]
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<text x="60" y="48" font-family="Arial" font-size="26" font-weight="700">Study workflow: reproducibility and external transportability audit</text>',
        '<text x="60" y="76" font-family="Arial" font-size="14" fill="#4B5563">Published LUAD transcriptomic signatures were separated by formula reproducibility, cutoff portability, and cross-cohort score transportability.</text>',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L8,3 z" fill="#374151"/></marker></defs>',
    ]
    for idx, ((x, y, w, h, title, sub), fill) in enumerate(zip(boxes, colors), start=1):
        parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="8" fill="{fill}" stroke="#9CA3AF"/>')
        parts.append(f'<text x="{x + 16}" y="{y + 30}" font-family="Arial" font-size="15" font-weight="700" fill="#111827">{idx}. {title}</text>')
        words = sub.split()
        line1 = " ".join(words[:5])
        line2 = " ".join(words[5:])
        parts.append(f'<text x="{x + 16}" y="{y + 56}" font-family="Arial" font-size="12" fill="#374151">{line1}</text>')
        if line2:
            parts.append(f'<text x="{x + 16}" y="{y + 73}" font-family="Arial" font-size="12" fill="#374151">{line2}</text>')
    for (x1, y1), (x2, y2) in arrows:
        parts.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#374151" stroke-width="2" marker-end="url(#arrow)"/>')
    parts.append('<text x="60" y="680" font-family="Arial" font-size="12" fill="#4B5563">Key reporting rule: median fallback analyses are score-portability analyses, not validation of original published risk-group thresholds.</text>')
    parts.append("</svg>")
    path.write_text("\n".join(parts), encoding="utf-8")


def main():
    PKG.mkdir(parents=True, exist_ok=True)
    (PKG / "tables").mkdir(exist_ok=True)
    (PKG / "figures").mkdir(exist_ok=True)
    (PKG / "supplementary").mkdir(exist_ok=True)
    (PKG / "notes").mkdir(exist_ok=True)

    selection = pd.read_csv(INPUTS["selection"])
    integrated = pd.read_csv(INPUTS["integrated"])
    geo_pairwise = pd.read_csv(INPUTS["geo_pairwise"])
    geo_signature = pd.read_csv(INPUTS["geo_signature"])
    robustness_summary = pd.read_csv(INPUTS["robustness_summary"])
    robustness_adjusted = pd.read_csv(INPUTS["robustness_adjusted"])
    robustness_zscore = pd.read_csv(INPUTS["robustness_zscore"])
    tcga_results = pd.read_csv(INPUTS["tcga_results"])
    cohort_source = pd.read_csv(INPUTS["cohort_source"])
    loco = pd.read_csv(INPUTS["loco"])

    t1 = table1(selection)
    t2 = integrated.copy()
    t1.to_csv(PKG / "tables" / "Table1_reproducible_signature_inventory.csv", index=False, encoding="utf-8")
    t2.to_csv(PKG / "tables" / "Table2_integrated_main_candidate_evidence.csv", index=False, encoding="utf-8")
    write_xlsx(PKG / "tables" / "Main_Tables.xlsx", [("Table1_Inventory", t1), ("Table2_Integrated", t2)])

    write_xlsx(
        PKG / "supplementary" / "Supplementary_Tables.xlsx",
        [
            ("S1_GEO_pairwise", geo_pairwise),
            ("S2_GEO_signature", geo_signature),
            ("S3_robustness_summary", robustness_summary),
            ("S4_clinical_adjusted", robustness_adjusted),
            ("S5_gene_zscore", robustness_zscore),
            ("S6_TCGA_sanity", tcga_results),
            ("S7_GEO_cohort_sources", cohort_source),
            ("S8_leave_one_cohort_out", loco),
        ],
    )

    make_workflow_svg(PKG / "figures" / "Figure1_workflow.svg")
    shutil.copyfile(INPUTS["heatmap"], PKG / "figures" / "Figure2_GEO_HR_heatmap.svg")
    shutil.copyfile(INPUTS["forest"], PKG / "figures" / "Figure3_pooled_HR_forest.svg")

    for key in ["a097_note", "gse50081_note", "core_text", "full_draft", "integrated_note", "methods_note"]:
        src = INPUTS[key]
        if src.exists():
            shutil.copyfile(src, PKG / "notes" / src.name)

    manifest = [
        "# LUAD Manuscript Package Manifest",
        "",
        "Generated from current project outputs.",
        "",
        "## Main Tables",
        "",
        "- `tables/Main_Tables.xlsx`: Table 1 reproducible-signature inventory and Table 2 integrated main-candidate evidence.",
        "- `tables/Table1_reproducible_signature_inventory.csv`",
        "- `tables/Table2_integrated_main_candidate_evidence.csv`",
        "",
        "## Figures",
        "",
        "- `figures/Figure1_workflow.svg`: study workflow.",
        "- `figures/Figure2_GEO_HR_heatmap.svg`: main GEO high-vs-low HR heatmap.",
        "- `figures/Figure3_pooled_HR_forest.svg`: descriptive fixed-effect pooled HR forest plot.",
        "",
        "## Supplementary",
        "",
        "- `supplementary/Supplementary_Tables.xlsx`: full GEO pairwise results, signature evidence, robustness analyses, TCGA sanity checks, GEO cohort source table, and leave-one-cohort-out summaries.",
        "",
        "## Notes",
        "",
        "- Full working manuscript draft, core text draft, reproducibility note, and targeted audit notes for A097 and GSE50081.",
        "",
        "## Key Positioning",
        "",
        "Write this as a reproducibility and external transportability audit. A131 is the only strict reported-cutoff candidate; the other main candidates support score portability under cohort-specific median grouping.",
    ]
    (PKG / "README.md").write_text("\n".join(manifest), encoding="utf-8")

    zip_path = OUT_DIR / "LUAD_manuscript_package.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for file in PKG.rglob("*"):
            if file.is_file():
                zf.write(file, file.relative_to(PKG.parent))

    print(f"Wrote package: {PKG}")
    print(f"Wrote zip: {zip_path}")


if __name__ == "__main__":
    main()
