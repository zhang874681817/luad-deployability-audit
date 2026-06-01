import csv
import gzip
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
DATA = ROOT / "data" / "raw" / "tcga_luad_xena"
OUT_DIR = ROOT / "outputs" / "luad_signature_extraction"

STAR_TPM = DATA / "TCGA-LUAD.star_tpm.tsv.gz"
SURV = DATA / "LUAD_survival.txt"
HGNC = WORK / "hgnc_complete_set.txt"
PREP = WORK / "luad_reproduction_preparation.json"
SELECTION = WORK / "luad_geo_manuscript_ready_signature_selection.csv"

OUT_RESULTS = WORK / "tcga_luad_main_signature_gdc_star_tpm_sanity_results.csv"
OUT_SAMPLE = WORK / "tcga_luad_main_signature_gdc_star_tpm_sample_scores.csv"
OUT_XLSX = OUT_DIR / "LUAD_TCGA_GDC_STAR_TPM_main_signature_sanity_check.xlsx"
OUT_MD = OUT_DIR / "LUAD_TCGA_GDC_STAR_TPM_main_signature_sanity_check_notes.md"

MAIN_TIERS = {"Tier_1_strict_reported_cutoff", "Tier_2_primary_median_fallback"}


def p_norm_two_sided(z):
    return math.erfc(abs(z) / math.sqrt(2.0)) if math.isfinite(z) else np.nan


def p_chisq1(x):
    return math.erfc(math.sqrt(x / 2.0)) if math.isfinite(x) and x >= 0 else np.nan


def cox_univariate(time, event, x, scale_x=False):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    x = np.asarray(x, dtype=float)
    mask = np.isfinite(time) & np.isfinite(event) & np.isfinite(x)
    time, event, x = time[mask], event[mask], x[mask]
    if len(time) < 30 or event.sum() < 10 or np.nanstd(x) == 0:
        return {}
    if scale_x:
        x = (x - np.mean(x)) / np.std(x)
    beta = 0.0
    event_times = np.sort(np.unique(time[event == 1]))
    for _ in range(60):
        grad = hess = 0.0
        xb = beta * x
        expxb = np.exp(np.clip(xb, -50, 50))
        for t in event_times:
            deaths = (time == t) & (event == 1)
            risk = time >= t
            d = deaths.sum()
            if d == 0:
                continue
            s0 = expxb[risk].sum()
            s1 = (x[risk] * expxb[risk]).sum()
            s2 = ((x[risk] ** 2) * expxb[risk]).sum()
            if s0 <= 0:
                continue
            grad += x[deaths].sum() - d * s1 / s0
            hess -= d * (s2 / s0 - (s1 / s0) ** 2)
        if hess == 0 or not math.isfinite(hess):
            return {}
        step = grad / hess
        beta -= step
        if abs(step) < 1e-7:
            break
    xb = beta * x
    expxb = np.exp(np.clip(xb, -50, 50))
    hess = 0.0
    for t in event_times:
        deaths = (time == t) & (event == 1)
        risk = time >= t
        d = deaths.sum()
        s0 = expxb[risk].sum()
        s1 = (x[risk] * expxb[risk]).sum()
        s2 = ((x[risk] ** 2) * expxb[risk]).sum()
        if s0 > 0:
            hess -= d * (s2 / s0 - (s1 / s0) ** 2)
    if hess >= 0:
        return {}
    se = math.sqrt(-1 / hess)
    z = beta / se if se else np.nan
    return {
        "HR": math.exp(beta),
        "CI95_low": math.exp(beta - 1.96 * se),
        "CI95_high": math.exp(beta + 1.96 * se),
        "p": p_norm_two_sided(z),
    }


def c_index(time, event, score):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    score = np.asarray(score, dtype=float)
    n = conc = ties = 0
    for i in range(len(time)):
        if event[i] != 1 or not np.isfinite(score[i]):
            continue
        for j in range(len(time)):
            if time[i] >= time[j] or not np.isfinite(score[j]):
                continue
            n += 1
            if score[i] > score[j]:
                conc += 1
            elif score[i] == score[j]:
                ties += 1
    return (conc + 0.5 * ties) / n if n else np.nan


def logrank(time, event, group):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    group = np.asarray(group, dtype=int)
    if len(set(group)) < 2:
        return {}
    o1 = e1 = v1 = 0.0
    for t in np.sort(np.unique(time[event == 1])):
        risk = time >= t
        deaths = (time == t) & (event == 1)
        n = risk.sum()
        n1 = (risk & (group == 1)).sum()
        d = deaths.sum()
        d1 = (deaths & (group == 1)).sum()
        if n <= 1:
            continue
        e = d * n1 / n
        v = n1 * (n - n1) * d * (n - d) / (n * n * (n - 1))
        o1 += d1
        e1 += e
        v1 += v
    if v1 <= 0:
        return {}
    chisq = (o1 - e1) ** 2 / v1
    return {"logrank_p": p_chisq1(chisq)}


def load_signatures(main_sigs):
    payload = json.loads(PREP.read_text(encoding="utf-8"))
    sigs = {sig: [] for sig in main_sigs}
    for row in payload["signature_gene_coefficients"]:
        sig = row["Signature_ID"]
        if sig not in sigs:
            continue
        gene = row.get("Approved_symbol") or row.get("Gene_original")
        coef = row.get("Coefficient_numeric")
        if gene and coef is not None:
            sigs[sig].append((gene, float(coef)))
    return sigs


def load_hgnc_map():
    h = pd.read_csv(HGNC, sep="\t", dtype=str)
    symbol_to_ensg = {}
    for _, row in h.iterrows():
        symbol = str(row.get("symbol", "")).strip()
        ensg = str(row.get("ensembl_gene_id", "")).strip()
        if symbol and ensg and ensg != "nan":
            symbol_to_ensg[symbol] = ensg
    return symbol_to_ensg


def sample_type(sample):
    parts = str(sample).split("-")
    return parts[3][:2] if len(parts) >= 4 else ""


def read_star_subset(ensg_ids):
    wanted = set(ensg_ids)
    rows = {}
    with gzip.open(STAR_TPM, "rt", encoding="utf-8", errors="replace") as f:
        header = f.readline().rstrip("\n").split("\t")
        samples = header[1:]
        tumor_idx = [i for i, s in enumerate(samples) if sample_type(s) == "01"]
        tumor_samples = [samples[i] for i in tumor_idx]
        for line in f:
            parts = line.rstrip("\n").split("\t")
            base_id = parts[0].split(".")[0]
            if base_id in wanted:
                rows[base_id] = [parts[i + 1] for i in tumor_idx]
    return pd.DataFrame(rows, index=tumor_samples).T.apply(pd.to_numeric, errors="coerce")


def load_survival(samples):
    surv = pd.read_csv(SURV, sep="\t")
    surv["patient"] = surv["sample"].astype(str).str[:12]
    surv["OS_event"] = pd.to_numeric(surv["OS"], errors="coerce")
    surv["OS_time_months"] = pd.to_numeric(surv["OS.time"], errors="coerce") / 30.4375
    sample_df = pd.DataFrame({"sample": samples, "patient": [s[:12] for s in samples]})
    out = sample_df.merge(surv[["patient", "OS_event", "OS_time_months"]], on="patient", how="inner")
    return out.dropna(subset=["OS_event", "OS_time_months"]).drop_duplicates("sample")


def write_xlsx(results, samples):
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in [("TCGA_GDC_STAR_TPM_Results", results), ("Sample_Scores", samples)]:
        ws = wb.create_sheet(name[:31])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            ws.append(list(row))
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = 34 if "Signature" in col or "genes" in col else 17
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    wb.save(OUT_XLSX)


def main():
    selection = pd.read_csv(SELECTION)
    main_rows = selection[selection["Final_analysis_tier"].isin(MAIN_TIERS)].copy()
    main_sigs = list(main_rows["Signature_ID"])
    meta = main_rows.set_index("Signature_ID").to_dict("index")
    sigs = load_signatures(main_sigs)
    hgnc = load_hgnc_map()

    needed = []
    for genes in sigs.values():
        for gene, _ in genes:
            if gene in hgnc:
                needed.append(hgnc[gene])
    expr = read_star_subset(needed)
    ensg_to_symbol = {v: k for k, v in hgnc.items()}
    expr_symbol = pd.DataFrame({ensg_to_symbol[i]: expr.loc[i] for i in expr.index if i in ensg_to_symbol}).T
    clin = load_survival(list(expr_symbol.columns))
    expr_symbol = expr_symbol[clin["sample"]]

    result_rows = []
    sample_rows = []
    for sig in main_sigs:
        genes = sigs[sig]
        missing = [g for g, _ in genes if g not in expr_symbol.index]
        base = {
            "Article_ID": meta[sig]["Article_ID"],
            "Signature_ID": sig,
            "Final_analysis_tier": meta[sig]["Final_analysis_tier"],
            "TCGA_expression_source": "GDC_STAR_TPM",
            "Gene_count": len(genes),
            "Present_gene_count": len(genes) - len(missing),
            "Missing_gene_count": len(missing),
            "Missing_genes": "; ".join(missing),
        }
        if missing:
            result_rows.append({**base, "status": "missing_genes"})
            continue
        coef = pd.Series({g: c for g, c in genes}, dtype=float)
        score = expr_symbol.loc[list(coef.index)].T.dot(coef)
        df = clin.copy()
        df["Risk_score"] = df["sample"].map(score)
        df = df.dropna(subset=["Risk_score"])
        cox_cont = cox_univariate(df["OS_time_months"], df["OS_event"], df["Risk_score"], scale_x=True)
        cutoff = float(df["Risk_score"].median())
        group = (df["Risk_score"] > cutoff).astype(int)
        cox_group = cox_univariate(df["OS_time_months"], df["OS_event"], group, scale_x=False)
        lr = logrank(df["OS_time_months"], df["OS_event"], group)
        result_rows.append(
            {
                **base,
                "status": "ok",
                "TCGA_N": len(df),
                "TCGA_events": int(df["OS_event"].sum()),
                "Score_median": cutoff,
                "C_index": c_index(df["OS_time_months"], df["OS_event"], df["Risk_score"]),
                "Continuous_HR_per_SD": cox_cont.get("HR", np.nan),
                "Continuous_p": cox_cont.get("p", np.nan),
                "Cutoff_rule_used": "cohort_median",
                "High_n": int(group.sum()),
                "Low_n": int((group == 0).sum()),
                "Group_HR": cox_group.get("HR", np.nan),
                "Group_p": cox_group.get("p", np.nan),
                "Logrank_p": lr.get("logrank_p", np.nan),
            }
        )
        for r in df.itertuples(index=False):
            sample_rows.append(
                {
                    "Article_ID": meta[sig]["Article_ID"],
                    "Signature_ID": sig,
                    "sample": r.sample,
                    "OS_time_months": r.OS_time_months,
                    "OS_event": r.OS_event,
                    "Risk_score": r.Risk_score,
                }
            )
    results = pd.DataFrame(result_rows)
    samples = pd.DataFrame(sample_rows)
    results.to_csv(OUT_RESULTS, index=False, encoding="utf-8")
    samples.to_csv(OUT_SAMPLE, index=False, encoding="utf-8")
    write_xlsx(results, samples)
    lines = [
        "# TCGA-LUAD GDC STAR TPM sanity check",
        "",
        "This analysis uses GDC harmonized TCGA-LUAD STAR TPM data from UCSC Xena. It was added because the legacy TCGA HiSeqV2 matrix lacks TMPRSS11E.",
        "",
    ]
    for _, row in results.iterrows():
        if row["status"] != "ok":
            lines.append(f"- {row['Article_ID']}: incomplete; missing {row['Missing_genes']}.")
        else:
            lines.append(
                f"- {row['Article_ID']}: N={int(row['TCGA_N'])}, events={int(row['TCGA_events'])}, "
                f"continuous HR/SD={row['Continuous_HR_per_SD']:.2f}, p={row['Continuous_p']:.3g}; "
                f"median-group HR={row['Group_HR']:.2f}, p={row['Group_p']:.3g}; C-index={row['C_index']:.3f}."
            )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_MD}")
    print(results[["Article_ID", "status", "Missing_genes", "Continuous_HR_per_SD", "Continuous_p", "Group_HR", "Group_p", "C_index"]].to_string(index=False))


if __name__ == "__main__":
    main()
