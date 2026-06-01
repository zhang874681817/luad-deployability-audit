import csv
import gzip
import json
import math
import re
import urllib.request
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
DATA = ROOT / "data" / "raw" / "tcga_luad_xena"
OUT_DIR = ROOT / "outputs" / "luad_signature_extraction"

PREP = WORK / "luad_reproduction_preparation.json"
SELECTION = WORK / "luad_geo_manuscript_ready_signature_selection.csv"

EXPR = DATA / "TCGA.LUAD.sampleMap.HiSeqV2.gz"
SURV = DATA / "LUAD_survival.txt"
CLINICAL = DATA / "LUAD_clinicalMatrix"

OUT_SAMPLE = WORK / "tcga_luad_main_signature_sample_scores.csv"
OUT_RESULTS = WORK / "tcga_luad_main_signature_sanity_results.csv"
OUT_XLSX = OUT_DIR / "LUAD_TCGA_main_signature_sanity_check.xlsx"
OUT_MD = OUT_DIR / "LUAD_TCGA_main_signature_sanity_check_notes.md"

MAIN_TIERS = {"Tier_1_strict_reported_cutoff", "Tier_2_primary_median_fallback"}
TCGA_SYMBOL_ALIASES = {
    "MELTF": "MFI2",
    "ERO1A": "ERO1L",
    "PKM": "PKM2",
}

URLS = {
    EXPR: "https://tcga.xenahubs.net/download/TCGA.LUAD.sampleMap/HiSeqV2.gz",
    SURV: "https://tcga.xenahubs.net/download/survival/LUAD_survival.txt",
    CLINICAL: "https://tcga.xenahubs.net/download/TCGA.LUAD.sampleMap/LUAD_clinicalMatrix",
}


def download(url, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.stat().st_size > 0:
        return
    tmp = path.with_suffix(path.suffix + ".tmp")
    print(f"Downloading {url}")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as r, tmp.open("wb") as out:
        while True:
            chunk = r.read(1024 * 1024)
            if not chunk:
                break
            out.write(chunk)
    tmp.replace(path)


def normal_p_two_sided(z):
    return math.erfc(abs(z) / math.sqrt(2.0)) if math.isfinite(z) else np.nan


def p_chisq1(x):
    if not math.isfinite(x) or x < 0:
        return np.nan
    return math.erfc(math.sqrt(x / 2.0))


def cox_univariate(time, event, x, scale_x=False):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    x = np.asarray(x, dtype=float)
    mask = np.isfinite(time) & np.isfinite(event) & np.isfinite(x)
    time, event, x = time[mask], event[mask], x[mask]
    if len(time) < 30 or event.sum() < 10 or np.nanstd(x) == 0:
        return {}
    if scale_x:
        x = (x - np.mean(x)) / np.std(x)
    beta = 0.0
    event_times = np.sort(np.unique(time[event == 1]))
    for _ in range(60):
        grad = 0.0
        hess = 0.0
        xb = beta * x
        expxb = np.exp(np.clip(xb, -50, 50))
        for t in event_times:
            deaths = (time == t) & (event == 1)
            risk = time >= t
            d = deaths.sum()
            if d == 0:
                continue
            s0 = expxb[risk].sum()
            s1 = (x[risk] * expxb[risk]).sum()
            s2 = ((x[risk] ** 2) * expxb[risk]).sum()
            if s0 <= 0:
                continue
            grad += x[deaths].sum() - d * s1 / s0
            hess -= d * (s2 / s0 - (s1 / s0) ** 2)
        if hess == 0 or not math.isfinite(hess):
            return {}
        step = grad / hess
        beta -= step
        if abs(step) < 1e-7:
            break
    xb = beta * x
    expxb = np.exp(np.clip(xb, -50, 50))
    hess = 0.0
    for t in event_times:
        deaths = (time == t) & (event == 1)
        risk = time >= t
        d = deaths.sum()
        s0 = expxb[risk].sum()
        s1 = (x[risk] * expxb[risk]).sum()
        s2 = ((x[risk] ** 2) * expxb[risk]).sum()
        if s0 > 0:
            hess -= d * (s2 / s0 - (s1 / s0) ** 2)
    if hess >= 0:
        return {}
    se = math.sqrt(-1.0 / hess)
    z = beta / se if se else np.nan
    return {
        "beta": beta,
        "HR": math.exp(beta),
        "CI95_low": math.exp(beta - 1.96 * se),
        "CI95_high": math.exp(beta + 1.96 * se),
        "p": normal_p_two_sided(z),
    }


def c_index(time, event, score):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    score = np.asarray(score, dtype=float)
    n = conc = ties = 0
    for i in range(len(time)):
        if event[i] != 1 or not np.isfinite(score[i]):
            continue
        for j in range(len(time)):
            if time[i] >= time[j] or not np.isfinite(score[j]):
                continue
            n += 1
            if score[i] > score[j]:
                conc += 1
            elif score[i] == score[j]:
                ties += 1
    return (conc + 0.5 * ties) / n if n else np.nan


def logrank(time, event, group):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    group = np.asarray(group, dtype=int)
    mask = np.isfinite(time) & np.isfinite(event) & np.isfinite(group)
    time, event, group = time[mask], event[mask], group[mask]
    if len(set(group)) < 2 or event.sum() < 10:
        return {}
    o1 = e1 = v1 = 0.0
    for t in np.sort(np.unique(time[event == 1])):
        risk = time >= t
        deaths = (time == t) & (event == 1)
        n = risk.sum()
        n1 = (risk & (group == 1)).sum()
        d = deaths.sum()
        d1 = (deaths & (group == 1)).sum()
        if n <= 1:
            continue
        exp1 = d * n1 / n
        var1 = n1 * (n - n1) * d * (n - d) / (n * n * (n - 1))
        o1 += d1
        e1 += exp1
        v1 += var1
    if v1 <= 0:
        return {}
    chisq = (o1 - e1) ** 2 / v1
    return {"logrank_chisq": chisq, "logrank_p": p_chisq1(chisq)}


def parse_cutoff(value):
    s = str(value or "")
    m = re.search(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", s)
    if not m:
        return None
    return float(m.group(0))


def load_signatures(main_sigs):
    payload = json.loads(PREP.read_text(encoding="utf-8"))
    readiness = {r["Signature_ID"]: r for r in payload["signature_readiness"]}
    sigs = {sig: {"genes": [], "readiness": readiness.get(sig, {})} for sig in main_sigs}
    for row in payload["signature_gene_coefficients"]:
        sig = row["Signature_ID"]
        if sig not in sigs:
            continue
        gene = row.get("Approved_symbol") or row.get("Gene_original")
        coef = row.get("Coefficient_numeric")
        if gene and coef is not None:
            sigs[sig]["genes"].append((gene, float(coef)))
    return sigs


def barcode_sample_type(sample):
    parts = str(sample).split("-")
    if len(parts) < 4:
        return ""
    return parts[3][:2]


def read_expression_subset(needed_symbols):
    needed_symbols = set(needed_symbols)
    rows = []
    with gzip.open(EXPR, "rt", encoding="utf-8", errors="replace") as f:
        header = f.readline().rstrip("\n").split("\t")
        samples = header[1:]
        tumor_keep = [i for i, s in enumerate(samples) if barcode_sample_type(s) == "01"]
        tumor_samples = [samples[i] for i in tumor_keep]
        for line in f:
            parts = line.rstrip("\n").split("\t")
            symbol = parts[0]
            if symbol not in needed_symbols:
                continue
            vals = [parts[i + 1] for i in tumor_keep]
            rows.append((symbol, vals))
    df = pd.DataFrame([vals for _, vals in rows], index=[symbol for symbol, _ in rows], columns=tumor_samples)
    return df.apply(pd.to_numeric, errors="coerce")


def load_tcga_survival():
    surv = pd.read_csv(SURV, sep="\t")
    lower = {c.lower(): c for c in surv.columns}
    sample_col = lower.get("sample") or lower.get("sampleid") or surv.columns[0]
    os_col = lower.get("os")
    time_col = lower.get("os.time") or lower.get("os_time") or lower.get("ostime")
    if os_col is None or time_col is None:
        raise RuntimeError(f"Could not identify OS columns in survival file: {list(surv.columns)}")
    out = surv[[sample_col, os_col, time_col]].copy()
    out.columns = ["sample", "OS_event", "OS_time_days"]
    out["patient"] = out["sample"].astype(str).str.slice(0, 12)
    out["OS_event"] = pd.to_numeric(out["OS_event"], errors="coerce")
    out["OS_time_months"] = pd.to_numeric(out["OS_time_days"], errors="coerce") / 30.4375
    out = out.dropna(subset=["patient", "OS_event", "OS_time_months"])
    out = out.sort_values("OS_time_months").drop_duplicates("patient", keep="last")
    return out


def score_signature(expr_symbol, genes):
    needed = [g for g, _ in genes]
    aliases_used = []
    resolved = []
    missing = []
    for g in needed:
        if g in expr_symbol.index:
            resolved.append(g)
        elif TCGA_SYMBOL_ALIASES.get(g) in expr_symbol.index:
            alias = TCGA_SYMBOL_ALIASES[g]
            resolved.append(alias)
            aliases_used.append(f"{g}->{alias}")
        else:
            missing.append(g)
    if missing:
        return None, missing, aliases_used
    coefs = pd.Series({resolved[i]: genes[i][1] for i in range(len(genes))}, dtype=float)
    return expr_symbol.loc[resolved].T.dot(coefs.loc[resolved]), [], aliases_used


def write_workbook(results, scores):
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in [("TCGA_Sanity_Results", results), ("TCGA_Sample_Scores", scores)]:
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            ws.append(list(row))
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = 38 if "Signature" in col or "missing" in col or "Title" in col else 17
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    wb.save(OUT_XLSX)


def main():
    for path, url in URLS.items():
        download(url, path)

    selection = pd.read_csv(SELECTION)
    main = selection[selection["Final_analysis_tier"].isin(MAIN_TIERS)].copy()
    main_sigs = list(main["Signature_ID"])
    sig_meta = main.set_index("Signature_ID").to_dict("index")
    sigs = load_signatures(main_sigs)

    needed_symbols = []
    for sig in main_sigs:
        for symbol, _ in sigs[sig]["genes"]:
            needed_symbols.append(symbol)
            if symbol in TCGA_SYMBOL_ALIASES:
                needed_symbols.append(TCGA_SYMBOL_ALIASES[symbol])
    expr_symbol = read_expression_subset(needed_symbols)
    expr_symbol = expr_symbol.loc[~expr_symbol.index.duplicated(keep="first")]

    surv = load_tcga_survival()
    sample_patient = pd.DataFrame({"sample": expr_symbol.columns, "patient": [s[:12] for s in expr_symbol.columns]})
    clin = sample_patient.merge(surv[["patient", "OS_event", "OS_time_months"]], on="patient", how="inner")
    expr_symbol = expr_symbol[clin["sample"]]

    result_rows = []
    sample_rows = []
    for sig in main_sigs:
        score, missing, aliases_used = score_signature(expr_symbol, sigs[sig]["genes"])
        meta = sig_meta[sig]
        base = {
            "Article_ID": meta["Article_ID"],
            "Signature_ID": sig,
            "Final_analysis_tier": meta["Final_analysis_tier"],
            "PMID": meta.get("PMID", ""),
            "Gene_count": len(sigs[sig]["genes"]),
            "Present_gene_count": len(sigs[sig]["genes"]) - len(missing),
            "Missing_gene_count": len(missing),
            "Missing_genes": "; ".join(missing),
            "TCGA_symbol_aliases_used": "; ".join(aliases_used),
            "Cutoff_type": meta.get("Cutoff_type", ""),
            "Cutoff_value": meta.get("Cutoff_value", ""),
        }
        if score is None:
            result_rows.append({**base, "status": "missing_genes"})
            continue
        df = clin.copy()
        df["Risk_score"] = df["sample"].map(score)
        df = df.dropna(subset=["Risk_score", "OS_event", "OS_time_months"])
        cox_cont = cox_univariate(df["OS_time_months"], df["OS_event"], df["Risk_score"], scale_x=True)
        cidx = c_index(df["OS_time_months"], df["OS_event"], df["Risk_score"])
        cutoff_numeric = parse_cutoff(meta.get("Cutoff_value", ""))
        cutoffs = [("cohort_median", float(df["Risk_score"].median()))]
        if cutoff_numeric is not None:
            cutoffs.insert(0, ("reported_numeric_cutoff", cutoff_numeric))
        first = True
        for cutoff_rule, cutoff in cutoffs:
            group = (df["Risk_score"] > cutoff).astype(int)
            high_n = int(group.sum())
            low_n = int((group == 0).sum())
            cox_group = cox_univariate(df["OS_time_months"], df["OS_event"], group, scale_x=False)
            lr = logrank(df["OS_time_months"], df["OS_event"], group)
            result_rows.append(
                {
                    **base,
                    "status": "ok" if high_n > 0 and low_n > 0 else "single_group_after_cutoff",
                    "TCGA_N": len(df),
                    "TCGA_events": int(df["OS_event"].sum()),
                    "Score_mean": float(df["Risk_score"].mean()),
                    "Score_sd": float(df["Risk_score"].std()),
                    "Score_median": float(df["Risk_score"].median()),
                    "C_index": cidx,
                    "Continuous_HR_per_SD": cox_cont.get("HR", np.nan),
                    "Continuous_CI95_low": cox_cont.get("CI95_low", np.nan),
                    "Continuous_CI95_high": cox_cont.get("CI95_high", np.nan),
                    "Continuous_p": cox_cont.get("p", np.nan),
                    "Cutoff_rule_used": cutoff_rule,
                    "Cutoff_used": cutoff,
                    "High_n": high_n,
                    "Low_n": low_n,
                    "Group_HR": cox_group.get("HR", np.nan),
                    "Group_CI95_low": cox_group.get("CI95_low", np.nan),
                    "Group_CI95_high": cox_group.get("CI95_high", np.nan),
                    "Group_p": cox_group.get("p", np.nan),
                    "Logrank_p": lr.get("logrank_p", np.nan),
                }
            )
            if first:
                for r in df.itertuples(index=False):
                    sample_rows.append(
                        {
                            "Article_ID": meta["Article_ID"],
                            "Signature_ID": sig,
                            "sample": r.sample,
                            "patient": r.patient,
                            "OS_time_months": r.OS_time_months,
                            "OS_event": r.OS_event,
                            "Risk_score": r.Risk_score,
                        }
                    )
                first = False

    results = pd.DataFrame(result_rows)
    samples = pd.DataFrame(sample_rows)
    results.to_csv(OUT_RESULTS, index=False, encoding="utf-8")
    samples.to_csv(OUT_SAMPLE, index=False, encoding="utf-8")
    write_workbook(results, samples)

    ok_median = results[(results["status"].eq("ok")) & (results["Cutoff_rule_used"].eq("cohort_median"))].copy()
    strict = results[(results["status"].eq("ok")) & (results["Cutoff_rule_used"].eq("reported_numeric_cutoff"))].copy()
    lines = [
        "# TCGA-LUAD main-signature sanity check",
        "",
        "Source: UCSC Xena TCGA-LUAD HiSeqV2 RNA-seq expression and LUAD survival table.",
        "",
        f"Primary tumor samples with OS data: {int(ok_median['TCGA_N'].max()) if len(ok_median) else 'NA'}.",
        "",
        "## Median-cutoff sanity check",
        "",
    ]
    for _, row in ok_median.iterrows():
        lines.append(
            f"- {row['Article_ID']}: continuous HR per SD {row['Continuous_HR_per_SD']:.2f}, p={row['Continuous_p']:.3g}; "
            f"median-group HR {row['Group_HR']:.2f}, p={row['Group_p']:.3g}; C-index {row['C_index']:.3f}."
        )
    if len(strict):
        lines.extend(["", "## Reported numeric cutoff check", ""])
        for _, row in strict.iterrows():
            lines.append(
                f"- {row['Article_ID']}: cutoff {row['Cutoff_used']}; high/low {int(row['High_n'])}/{int(row['Low_n'])}; "
                f"group HR {row['Group_HR']:.2f}, p={row['Group_p']:.3g}."
            )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_MD}")
    print(f"Results rows: {len(results)}")
    print(f"Sample score rows: {len(samples)}")


if __name__ == "__main__":
    main()
