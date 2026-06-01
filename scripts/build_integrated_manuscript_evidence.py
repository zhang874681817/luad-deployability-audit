from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
OUT_DIR = ROOT / "outputs" / "luad_signature_extraction"

GEO_SIG = WORK / "luad_geo_manuscript_signature_evidence.csv"
ROBUST = WORK / "luad_geo_main_signature_robustness_summary.csv"
TCGA = WORK / "tcga_luad_main_signature_gdc_star_tpm_sanity_results.csv"
DECISION = WORK / "luad_geo_manuscript_decision_ranking.csv"

OUT_CSV = WORK / "integrated_main_signature_manuscript_evidence.csv"
OUT_XLSX = OUT_DIR / "LUAD_integrated_main_signature_manuscript_evidence.xlsx"
OUT_MD = OUT_DIR / "LUAD_integrated_manuscript_recommendation.md"


def num(s):
    return pd.to_numeric(s, errors="coerce")


def fmt(v, digits=3):
    if pd.isna(v):
        return "NA"
    return f"{float(v):.{digits}f}"


def main_call(row):
    article = row["Article_ID"]
    if article == "A131":
        return "Best strict-cutoff candidate for main text"
    if article in {"A097", "A067", "A168", "A001", "A059"}:
        return "Strong score-portability candidate"
    if article == "A129":
        return "Strong score-portability candidate"
    if article in {"A098"}:
        return "Moderate score-portability candidate"
    if article in {"A065", "A173"}:
        return "Supportive but lower-priority candidate"
    return "Review manually"


def main():
    geo = pd.read_csv(GEO_SIG)
    robust = pd.read_csv(ROBUST)
    decision = pd.read_csv(DECISION)
    tcga = pd.read_csv(TCGA)

    main_geo = geo[geo["Main_or_sensitivity"].eq("main")].copy()
    robust = robust.copy()
    tcga_median = tcga[(tcga["status"].eq("ok")) & (tcga["Cutoff_rule_used"].eq("cohort_median"))].copy()
    tcga_strict = tcga[(tcga["status"].eq("ok")) & (tcga["Cutoff_rule_used"].eq("reported_numeric_cutoff"))].copy()
    tcga_missing = tcga[tcga["status"].eq("missing_genes")].copy()

    out = main_geo.merge(robust, on=["Article_ID", "Signature_ID", "Final_analysis_tier"], how="left", suffixes=("", "_robust"))
    tcga_cols = [
        "Article_ID",
        "TCGA_N",
        "TCGA_events",
        "C_index",
        "Continuous_HR_per_SD",
        "Continuous_p",
        "Group_HR",
        "Group_p",
        "TCGA_expression_source",
    ]
    out = out.merge(
        tcga_median[tcga_cols].rename(
            columns={
                "C_index": "TCGA_median_C_index",
                "Continuous_HR_per_SD": "TCGA_continuous_HR_per_SD",
                "Continuous_p": "TCGA_continuous_p",
                "Group_HR": "TCGA_median_group_HR",
                "Group_p": "TCGA_median_group_p",
            }
        ),
        on="Article_ID",
        how="left",
    )
    if {"Cutoff_used", "High_n", "Low_n", "Group_HR", "Group_p"}.issubset(tcga_strict.columns):
        out = out.merge(
            tcga_strict[["Article_ID", "Cutoff_used", "High_n", "Low_n", "Group_HR", "Group_p"]].rename(
                columns={
                    "Cutoff_used": "TCGA_reported_cutoff_used",
                    "High_n": "TCGA_reported_cutoff_high_n",
                    "Low_n": "TCGA_reported_cutoff_low_n",
                    "Group_HR": "TCGA_reported_cutoff_group_HR",
                    "Group_p": "TCGA_reported_cutoff_group_p",
                }
            ),
            on="Article_ID",
            how="left",
        )
    else:
        for col in [
            "TCGA_reported_cutoff_used",
            "TCGA_reported_cutoff_high_n",
            "TCGA_reported_cutoff_low_n",
            "TCGA_reported_cutoff_group_HR",
            "TCGA_reported_cutoff_group_p",
        ]:
            out[col] = np.nan
    out = out.merge(
        tcga_missing[["Article_ID", "Missing_gene_count", "Missing_genes"]].rename(
            columns={
                "Missing_gene_count": "TCGA_missing_gene_count",
                "Missing_genes": "TCGA_missing_genes",
            }
        ),
        on="Article_ID",
        how="left",
    )

    out["Manuscript_priority_call"] = out.apply(main_call, axis=1)
    out["Evidence_summary"] = out.apply(
        lambda r: (
            f"GEO grouped {int(r['Group_sig_p05_n'])}/{int(r['GEO_tests_ok'])}; "
            f"GEO continuous {int(r['Continuous_sig_p05_n'])}/{int(r['GEO_tests_ok'])}; "
            f"adjusted {int(r['Adjusted_risk_p05_n'])}/{int(r['Adjusted_models_n'])}; "
            f"median Delta C-index vs clinical {fmt(r['Median_delta_C_index_vs_clinical'])}; "
            f"zscore grouped {int(r['Zscore_group_p05_n'])}/{int(r['Zscore_tests_n'])}; "
            f"TCGA median HR {fmt(r['TCGA_median_group_HR'])}, p {fmt(r['TCGA_median_group_p'])}."
        ),
        axis=1,
    )
    out["Critical_limitations"] = out.apply(
        lambda r: (
            "A097 published TCGA cutoff is not portable to GEO processed matrices; present as score portability."
            if r["Article_ID"] == "A097"
            else (
                "Reported cutoff is the only strict cross-platform grouping candidate."
                if r["Article_ID"] == "A131"
                else "Grouped GEO results use cohort-median fallback; not original cutoff validation."
            )
        ),
        axis=1,
    )

    priority_order = {
        "Best strict-cutoff candidate for main text": 1,
        "Strong score-portability candidate": 2,
        "Moderate score-portability candidate": 3,
        "Supportive but lower-priority candidate": 4,
    }
    out["_priority"] = out["Manuscript_priority_call"].map(priority_order).fillna(99)
    out = out.sort_values(["_priority", "Group_sig_p05_n", "Adjusted_risk_p05_n", "TCGA_median_C_index"], ascending=[True, False, False, False])

    keep = [
        "Manuscript_priority_call",
        "Article_ID",
        "Signature_ID",
        "Final_analysis_tier",
        "GEO_tests_ok",
        "Group_sig_p05_n",
        "Group_sig_BH_main_q05_n",
        "Continuous_sig_p05_n",
        "Continuous_sig_BH_main_q05_n",
        "Median_C_index",
        "Median_group_HR",
        "Pooled_HR_fixed",
        "Pooled_HR_CI95_low",
        "Pooled_HR_CI95_high",
        "Meta_I2_percent",
        "Adjusted_models_n",
        "Adjusted_risk_p05_n",
        "Adjusted_median_HR",
        "Clinical_benchmark_models_n",
        "Median_delta_C_index_vs_clinical",
        "Delta_C_index_gt0_n",
        "Zscore_tests_n",
        "Zscore_group_p05_n",
        "Zscore_continuous_p05_n",
        "Zscore_median_C_index",
        "TCGA_N",
        "TCGA_events",
        "TCGA_median_C_index",
        "TCGA_continuous_HR_per_SD",
        "TCGA_continuous_p",
        "TCGA_median_group_HR",
        "TCGA_median_group_p",
        "TCGA_reported_cutoff_used",
        "TCGA_reported_cutoff_high_n",
        "TCGA_reported_cutoff_low_n",
        "TCGA_reported_cutoff_group_HR",
        "TCGA_reported_cutoff_group_p",
        "TCGA_expression_source",
        "TCGA_missing_gene_count",
        "TCGA_missing_genes",
        "Evidence_summary",
        "Critical_limitations",
    ]
    out = out[[c for c in keep if c in out.columns]].copy()
    out.to_csv(OUT_CSV, index=False, encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Integrated_Evidence")
    ws.append(list(out.columns))
    for row in out.itertuples(index=False, name=None):
        ws.append(list(row))
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(out.columns, start=1):
        width = 48 if any(x in col for x in ["Evidence", "limitations", "Signature", "priority"]) else 17
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, float):
                cell.number_format = "0.000"

    ws2 = wb.create_sheet("Reviewer_Positioning")
    notes = [
        ("Recommended manuscript frame", "Reproducibility and transportability audit of published LUAD transcriptomic prognostic signatures."),
        ("Main strict result", "A131 is the only current signature suitable for strict reported-cutoff cross-platform grouping claims."),
        ("Main score-portability result", "A129, A067, A097, A168, A001, and A059 are strong score-portability candidates. A129 required GDC STAR TPM rather than legacy HiSeqV2 because TMPRSS11E is absent from the legacy matrix."),
        ("Cutoff caveat", "Median-fallback grouping must be described as cohort-specific score portability, not validation of original clinical cutoffs."),
        ("Clinical added value", "Clinical-plus-risk C-index exceeded clinical-only C-index in all benchmarkable main-candidate models in the current robustness screen."),
        ("Remaining high-impact gap", "Final manuscript should add methods-level transparency on multiple testing, cohort reuse, and TCGA-as-training sanity check rather than independent validation."),
    ]
    ws2.append(["Topic", "Recommendation"])
    for row in notes:
        ws2.append(row)
    for cell in ws2[1]:
        cell.fill = PatternFill("solid", fgColor="7030A0")
        cell.font = Font(color="FFFFFF", bold=True)
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 110
    for row in ws2.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT_XLSX)

    lines = [
        "# Integrated manuscript recommendation",
        "",
        "Recommended framing: this should be written as a reproducibility and external transportability audit of published LUAD transcriptomic prognostic signatures, not as a new-signature discovery paper.",
        "",
        "## Main claims now supported",
        "",
        "- A131 is the strongest strict-cutoff candidate. It remains significant in GEO using the reported cutoff and also passes TCGA sanity using the reported numeric cutoff.",
        "- A129, A067, A097, A168, A001, and A059 show strong score portability across GEO, clinical adjustment, z-score sensitivity, and TCGA sanity checks.",
        "- A129 required the GDC STAR TPM TCGA-LUAD matrix for complete formula rescoring because TMPRSS11E is absent from the legacy TCGA HiSeqV2 matrix.",
        "- A098 is moderate but supported; A065 and A173 are lower-priority supporting candidates.",
        "",
        "## Manuscript caution",
        "",
        "Median-fallback grouped results should be described as cohort-specific score-portability analyses. They should not be phrased as validation of the original published risk-group thresholds.",
        "",
        "TCGA-LUAD is primarily a formula-direction sanity check because many signatures were originally trained on TCGA-LUAD. It should not be counted as an independent external validation cohort for those signatures.",
        "",
        "## Priority table",
        "",
    ]
    for _, r in out.iterrows():
        lines.append(f"- {r['Article_ID']}: {r['Manuscript_priority_call']}. {r['Evidence_summary']} Limitation: {r['Critical_limitations']}")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_MD}")
    print(f"Wrote {OUT_CSV}")


if __name__ == "__main__":
    main()
