import json
import math
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
OUT_DIR = ROOT / "outputs" / "luad_signature_extraction"

PREP = WORK / "luad_reproduction_preparation.json"
CLIN = WORK / "geo_clinical_samples.csv"
READY = WORK / "geo_expression_clinical_readiness.csv"
SELECTION = WORK / "luad_geo_manuscript_ready_signature_selection.csv"

OUT_ADJ = WORK / "luad_geo_main_signature_clinical_adjusted_cox.csv"
OUT_STD = WORK / "luad_geo_main_signature_zscore_sensitivity.csv"
OUT_SUMMARY = WORK / "luad_geo_main_signature_robustness_summary.csv"
OUT_XLSX = OUT_DIR / "LUAD_GEO_main_signature_robustness_pack.xlsx"
OUT_MD = OUT_DIR / "LUAD_GEO_main_signature_robustness_notes.md"

MAIN_TIERS = {"Tier_1_strict_reported_cutoff", "Tier_2_primary_median_fallback"}


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def split_cohorts(v):
    out = set()
    for token in re.split(r"[;/|,]", str(v or "")):
        t = norm(token).upper()
        if not t or t == "NAN":
            continue
        out.add(t.replace("TCGA LUAD", "TCGA-LUAD"))
    return out


def is_luad_like(row):
    text = " ".join(str(row.get(c, "")) for c in ["histology", "sample_type", "stage"]).lower()
    if any(x in text for x in ["normal", "squamous", "large cell", "ntl"]):
        return False
    return any(x in text for x in ["adeno", "adenocarcinoma", "lung tumour", "lung tumor", "primary lung tumor", "nsclc"])


def p_norm_two_sided(z):
    return math.erfc(abs(z) / math.sqrt(2.0)) if math.isfinite(z) else np.nan


def cox_ph(time, event, x, names):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    x = np.asarray(x, dtype=float)
    mask = np.isfinite(time) & np.isfinite(event) & np.isfinite(x).all(axis=1)
    time, event, x = time[mask], event[mask], x[mask]
    if len(time) < 30 or event.sum() < 10 or x.shape[1] == 0:
        return {}
    keep = [i for i in range(x.shape[1]) if np.nanstd(x[:, i]) > 0]
    if not keep:
        return {}
    x = x[:, keep]
    names = [names[i] for i in keep]
    p = x.shape[1]
    beta = np.zeros(p)
    event_times = np.sort(np.unique(time[event == 1]))
    for _ in range(60):
        grad = np.zeros(p)
        hess = np.zeros((p, p))
        xb = np.clip(x @ beta, -50, 50)
        expxb = np.exp(xb)
        for t in event_times:
            deaths = (time == t) & (event == 1)
            risk = time >= t
            d = deaths.sum()
            if d == 0:
                continue
            xr = x[risk]
            w = expxb[risk]
            s0 = w.sum()
            if s0 <= 0:
                continue
            s1 = (xr * w[:, None]).sum(axis=0)
            s2 = (xr[:, :, None] * xr[:, None, :] * w[:, None, None]).sum(axis=0)
            mean = s1 / s0
            grad += x[deaths].sum(axis=0) - d * mean
            hess -= d * (s2 / s0 - np.outer(mean, mean))
        try:
            step = np.linalg.solve(hess, grad)
        except np.linalg.LinAlgError:
            return {}
        beta -= step
        if np.max(np.abs(step)) < 1e-7:
            break
    xb = np.clip(x @ beta, -50, 50)
    expxb = np.exp(xb)
    hess = np.zeros((p, p))
    for t in event_times:
        deaths = (time == t) & (event == 1)
        risk = time >= t
        d = deaths.sum()
        xr = x[risk]
        w = expxb[risk]
        s0 = w.sum()
        if s0 <= 0:
            continue
        s1 = (xr * w[:, None]).sum(axis=0)
        s2 = (xr[:, :, None] * xr[:, None, :] * w[:, None, None]).sum(axis=0)
        mean = s1 / s0
        hess -= d * (s2 / s0 - np.outer(mean, mean))
    try:
        cov = np.linalg.inv(-hess)
    except np.linalg.LinAlgError:
        return {}
    out = {"N_model": int(len(time)), "Events_model": int(event.sum())}
    for i, name in enumerate(names):
        se = math.sqrt(cov[i, i]) if cov[i, i] > 0 else np.nan
        z = beta[i] / se if se and math.isfinite(se) else np.nan
        out[f"{name}_beta"] = beta[i]
        out[f"{name}_HR"] = math.exp(beta[i])
        out[f"{name}_CI95_low"] = math.exp(beta[i] - 1.96 * se) if math.isfinite(se) else np.nan
        out[f"{name}_CI95_high"] = math.exp(beta[i] + 1.96 * se) if math.isfinite(se) else np.nan
        out[f"{name}_p"] = p_norm_two_sided(z)
    return out


def c_index(time, event, score):
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    score = np.asarray(score, dtype=float)
    n = conc = ties = 0
    for i in range(len(time)):
        if event[i] != 1 or not np.isfinite(score[i]):
            continue
        for j in range(len(time)):
            if time[i] >= time[j] or not np.isfinite(score[j]):
                continue
            n += 1
            if score[i] > score[j]:
                conc += 1
            elif score[i] == score[j]:
                ties += 1
    return (conc + 0.5 * ties) / n if n else np.nan


def stage_major(value):
    s = str(value or "").upper().strip()
    if not s or s == "NAN":
        return np.nan
    if re.search(r"\bIV\b|^4", s):
        return 4.0
    if re.search(r"\bIII\b|^3", s):
        return 3.0
    if re.search(r"\bII\b|^2", s):
        return 2.0
    if re.search(r"\bI\b|^1", s):
        return 1.0
    return np.nan


def sex_binary(value):
    s = str(value or "").upper().strip()
    if s in {"M", "MALE", "1"}:
        return 1.0
    if s in {"F", "FEMALE", "0"}:
        return 0.0
    return np.nan


def z(v):
    v = pd.to_numeric(pd.Series(v), errors="coerce")
    sd = v.std(skipna=True)
    if not np.isfinite(sd) or sd == 0:
        return pd.Series([np.nan] * len(v))
    return (v - v.mean(skipna=True)) / sd


def lp_from_model(df, covariates, labels, model):
    lp = np.zeros(len(df), dtype=float)
    used = False
    for col, label in zip(covariates, labels):
        beta = model.get(f"{label}_beta")
        if beta is None or not np.isfinite(beta):
            continue
        lp += beta * pd.to_numeric(df[col], errors="coerce").to_numpy(float)
        used = True
    return lp if used else None


def load_signatures():
    payload = json.loads(PREP.read_text(encoding="utf-8"))
    genes = {}
    readiness = {r["Signature_ID"]: r for r in payload["signature_readiness"]}
    for row in payload["signature_gene_coefficients"]:
        sig = row["Signature_ID"]
        gene = row.get("Approved_symbol") or row.get("Gene_original")
        coef = row.get("Coefficient_numeric")
        if gene and coef is not None:
            genes.setdefault(sig, []).append((gene, float(coef)))
    return {sig: {"genes": g, "readiness": readiness.get(sig, {})} for sig, g in genes.items()}


def load_expression(path):
    return pd.read_csv(path, sep="\t").set_index("Gene_symbol")


def score_expression(expr, genes, transform):
    needed = [g for g, _ in genes]
    if not all(g in expr.index for g in needed):
        return None
    sub = expr.loc[needed].astype(float)
    if transform == "gene_zscore":
        arr = sub.to_numpy(float)
        means = np.nanmean(arr, axis=1)
        sds = np.nanstd(arr, axis=1)
        sds[sds == 0] = np.nan
        arr = (arr - means[:, None]) / sds[:, None]
        sub = pd.DataFrame(arr, index=sub.index, columns=sub.columns)
    coefs = pd.Series({g: c for g, c in genes}, dtype=float)
    return sub.T.dot(coefs.loc[sub.index])


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    selection = pd.read_csv(SELECTION)
    main_sigs = selection[selection["Final_analysis_tier"].isin(MAIN_TIERS)].copy()
    sig_order = list(main_sigs["Signature_ID"])
    sig_meta = main_sigs.set_index("Signature_ID").to_dict("index")
    signatures = load_signatures()

    clin = pd.read_csv(CLIN)
    clin["LUAD_like"] = clin.apply(is_luad_like, axis=1)
    clin["age_z"] = z(clin["age"])
    clin["sex_male"] = clin["sex"].apply(sex_binary)
    clin["stage_major"] = clin["stage"].apply(stage_major)
    clin["stage_z"] = z(clin["stage_major"])

    scores = pd.read_csv(WORK / "geo_external_validation_sample_scores.csv")
    scores = scores[scores["Signature_ID"].isin(sig_order)].copy()
    scores = scores.merge(
        clin[["GSE", "geo_accession", "LUAD_like", "age_z", "sex_male", "stage_z"]],
        on=["GSE", "geo_accession"],
        how="left",
    )
    adj_rows = []
    for (sig, gse), sub in scores.groupby(["Signature_ID", "GSE"]):
        sub = sub[sub["LUAD_like"].eq(True)].copy()
        sub["risk_z"] = z(sub["Risk_score"])
        covariates = ["risk_z"]
        cov_labels = ["risk_score_per_SD"]
        for col, label in [("age_z", "age_per_SD"), ("sex_male", "sex_male"), ("stage_z", "stage_per_SD")]:
            nonmiss = sub[col].notna().sum()
            if nonmiss >= 50 and sub[col].nunique(dropna=True) >= 2:
                covariates.append(col)
                cov_labels.append(label)
        model = cox_ph(sub["OS_time_months"], sub["OS_event"], sub[covariates].to_numpy(float), cov_labels)
        if not model or "risk_score_per_SD_HR" not in model:
            continue
        complete_mask = (
            pd.to_numeric(sub["OS_time_months"], errors="coerce").notna()
            & pd.to_numeric(sub["OS_event"], errors="coerce").notna()
            & sub[covariates].apply(pd.to_numeric, errors="coerce").notna().all(axis=1)
        )
        comp = sub[complete_mask].copy()
        clinical_covariates = covariates[1:]
        clinical_labels = cov_labels[1:]
        clinical_model = {}
        clinical_c = np.nan
        full_c = np.nan
        delta_c = np.nan
        risk_only_c = c_index(comp["OS_time_months"], comp["OS_event"], comp["risk_z"]) if len(comp) else np.nan
        if clinical_covariates:
            clinical_model = cox_ph(comp["OS_time_months"], comp["OS_event"], comp[clinical_covariates].to_numpy(float), clinical_labels)
            clin_lp = lp_from_model(comp, clinical_covariates, clinical_labels, clinical_model)
            full_lp = lp_from_model(comp, covariates, cov_labels, model)
            if clin_lp is not None:
                clinical_c = c_index(comp["OS_time_months"], comp["OS_event"], clin_lp)
            if full_lp is not None:
                full_c = c_index(comp["OS_time_months"], comp["OS_event"], full_lp)
            if np.isfinite(clinical_c) and np.isfinite(full_c):
                delta_c = full_c - clinical_c
        adj_rows.append(
            {
                "Article_ID": sig_meta[sig]["Article_ID"],
                "Signature_ID": sig,
                "Final_analysis_tier": sig_meta[sig]["Final_analysis_tier"],
                "GSE": gse,
                "Covariates_in_model": "; ".join(cov_labels),
                "Clinical_covariates_only": "; ".join(clinical_labels),
                "Risk_only_C_index": risk_only_c,
                "Clinical_only_C_index": clinical_c,
                "Risk_plus_clinical_C_index": full_c,
                "Delta_C_index_vs_clinical": delta_c,
                **model,
            }
        )
    adj = pd.DataFrame(adj_rows)
    adj.to_csv(OUT_ADJ, index=False, encoding="utf-8")

    ready = pd.read_csv(READY)
    ready = ready[ready["Machine_ready_LUAD_like"].eq("yes") & ready["Matrix_status"].eq("ok")]
    expr_cache = {}
    std_rows = []
    for _, cohort in ready.iterrows():
        gse = cohort["GSE"]
        expr_path = Path(str(cohort["Scoring_expression_file"]))
        if not expr_path.exists():
            continue
        expr = expr_cache.setdefault(gse, load_expression(expr_path))
        csub = clin[
            clin["GSE"].eq(gse)
            & clin["LUAD_like"].eq(True)
            & clin["geo_accession"].isin(expr.columns)
            & clin["OS_time_months"].notna()
            & clin["OS_event"].notna()
        ].copy()
        if len(csub) < 30:
            continue
        for sig in sig_order:
            train = split_cohorts(signatures[sig]["readiness"].get("Training_cohort", ""))
            if gse.upper() in train:
                continue
            score = score_expression(expr, signatures[sig]["genes"], "gene_zscore")
            if score is None:
                continue
            s = csub.merge(score.rename("Risk_score_zexpr"), left_on="geo_accession", right_index=True, how="left")
            s = s[s["Risk_score_zexpr"].notna()].copy()
            s["Risk_score_zexpr_SD"] = z(s["Risk_score_zexpr"])
            cox_cont = cox_ph(s["OS_time_months"], s["OS_event"], s[["Risk_score_zexpr_SD"]].to_numpy(float), ["risk_score_per_SD"])
            cutoff = s["Risk_score_zexpr"].median()
            s["high_group"] = (s["Risk_score_zexpr"] > cutoff).astype(float)
            cox_group = cox_ph(s["OS_time_months"], s["OS_event"], s[["high_group"]].to_numpy(float), ["high_vs_low"])
            std_rows.append(
                {
                    "Article_ID": sig_meta[sig]["Article_ID"],
                    "Signature_ID": sig,
                    "Final_analysis_tier": sig_meta[sig]["Final_analysis_tier"],
                    "GSE": gse,
                    "N": len(s),
                    "Events": int(pd.to_numeric(s["OS_event"], errors="coerce").sum()),
                    "Zexpr_C_index": c_index(s["OS_time_months"], s["OS_event"], s["Risk_score_zexpr"]),
                    "Zexpr_continuous_HR_per_SD": cox_cont.get("risk_score_per_SD_HR", np.nan),
                    "Zexpr_continuous_p": cox_cont.get("risk_score_per_SD_p", np.nan),
                    "Zexpr_group_HR": cox_group.get("high_vs_low_HR", np.nan),
                    "Zexpr_group_p": cox_group.get("high_vs_low_p", np.nan),
                    "Zexpr_cutoff_rule": "cohort_median_after_gene_zscore",
                }
            )
    std = pd.DataFrame(std_rows)
    std.to_csv(OUT_STD, index=False, encoding="utf-8")

    summary_rows = []
    for sig in sig_order:
        a = adj[adj["Signature_ID"].eq(sig)] if len(adj) else pd.DataFrame()
        s = std[std["Signature_ID"].eq(sig)] if len(std) else pd.DataFrame()
        summary_rows.append(
            {
                "Article_ID": sig_meta[sig]["Article_ID"],
                "Signature_ID": sig,
                "Final_analysis_tier": sig_meta[sig]["Final_analysis_tier"],
                "Adjusted_models_n": len(a),
                "Adjusted_risk_p05_n": int((pd.to_numeric(a.get("risk_score_per_SD_p", pd.Series(dtype=float)), errors="coerce") < 0.05).sum()) if len(a) else 0,
                "Adjusted_risk_HR_gt1_n": int((pd.to_numeric(a.get("risk_score_per_SD_HR", pd.Series(dtype=float)), errors="coerce") > 1).sum()) if len(a) else 0,
                "Adjusted_median_HR": pd.to_numeric(a.get("risk_score_per_SD_HR", pd.Series(dtype=float)), errors="coerce").median() if len(a) else np.nan,
                "Clinical_benchmark_models_n": int(pd.to_numeric(a.get("Clinical_only_C_index", pd.Series(dtype=float)), errors="coerce").notna().sum()) if len(a) else 0,
                "Median_delta_C_index_vs_clinical": pd.to_numeric(a.get("Delta_C_index_vs_clinical", pd.Series(dtype=float)), errors="coerce").median() if len(a) else np.nan,
                "Delta_C_index_gt0_n": int((pd.to_numeric(a.get("Delta_C_index_vs_clinical", pd.Series(dtype=float)), errors="coerce") > 0).sum()) if len(a) else 0,
                "Zscore_tests_n": len(s),
                "Zscore_group_p05_n": int((pd.to_numeric(s.get("Zexpr_group_p", pd.Series(dtype=float)), errors="coerce") < 0.05).sum()) if len(s) else 0,
                "Zscore_continuous_p05_n": int((pd.to_numeric(s.get("Zexpr_continuous_p", pd.Series(dtype=float)), errors="coerce") < 0.05).sum()) if len(s) else 0,
                "Zscore_group_HR_gt1_n": int((pd.to_numeric(s.get("Zexpr_group_HR", pd.Series(dtype=float)), errors="coerce") > 1).sum()) if len(s) else 0,
                "Zscore_median_C_index": pd.to_numeric(s.get("Zexpr_C_index", pd.Series(dtype=float)), errors="coerce").median() if len(s) else np.nan,
                "Zscore_median_group_HR": pd.to_numeric(s.get("Zexpr_group_HR", pd.Series(dtype=float)), errors="coerce").median() if len(s) else np.nan,
            }
        )
    summary = pd.DataFrame(summary_rows)
    summary.to_csv(OUT_SUMMARY, index=False, encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)
    for name, df in [("Summary", summary), ("Clinical_Adjusted_Cox", adj), ("Gene_Zscore_Sensitivity", std)]:
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            ws.append(list(row))
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = 34 if "Signature" in col or "Covariates" in col else 17
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    wb.save(OUT_XLSX)

    lines = [
        "# Main signature robustness notes",
        "",
        "Clinical adjustment used GEO-level covariates available in each cohort: risk score per SD was always included; age, sex, and stage were added when at least 50 samples had data and the covariate varied. When at least one clinical covariate was usable, a clinical-only Cox score was also compared with a clinical-plus-risk Cox score by C-index.",
        "",
        "Gene-zscore sensitivity rescored each signature after standardizing each signature gene within each GEO cohort. Grouped tests then used the cohort median of the z-scored signature score.",
        "",
        "## Summary",
        "",
    ]
    for _, r in summary.iterrows():
        lines.append(
            f"- {r['Article_ID']}: adjusted risk p<0.05 in {int(r['Adjusted_risk_p05_n'])}/{int(r['Adjusted_models_n'])}; "
            f"median Delta C-index vs clinical {r['Median_delta_C_index_vs_clinical']:.3f} across {int(r['Clinical_benchmark_models_n'])} benchmarkable models; "
            f"gene-zscore grouped p<0.05 in {int(r['Zscore_group_p05_n'])}/{int(r['Zscore_tests_n'])}; "
            f"gene-zscore continuous p<0.05 in {int(r['Zscore_continuous_p05_n'])}/{int(r['Zscore_tests_n'])}; "
            f"median zscore C-index {r['Zscore_median_C_index']:.3f}."
        )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_MD}")
    print(f"Adjusted models: {len(adj)}")
    print(f"Zscore sensitivity tests: {len(std)}")


if __name__ == "__main__":
    main()
