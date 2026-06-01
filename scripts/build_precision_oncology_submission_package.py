import html
import math
import re
import shutil
import textwrap
import zipfile
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_BASE = ROOT / "outputs" / "luad_signature_extraction"
AUDIT = OUT_BASE / "precision_oncology_deployability_audit"
FIG_DIR = AUDIT / "figures"
PKG_DIR = OUT_BASE / "precision_oncology_deployability_package"

XLSX = AUDIT / "LUAD_precision_oncology_deployability_audit_pack.xlsx"
MANUSCRIPT = AUDIT / "LUAD_precision_oncology_deployability_manuscript_draft.md"
DOCX = AUDIT / "LUAD_precision_oncology_deployability_manuscript_draft.docx"
ZIP = OUT_BASE / "LUAD_precision_oncology_deployability_package.zip"

PRIMARY_ARTICLE_ORDER = ["A131", "A129", "A067", "A097", "A168", "A001", "A059", "A098", "A065", "A173"]

COLORS = {
    "ink": "1F2937",
    "muted": "6B7280",
    "blue": "2F5597",
    "teal": "0F766E",
    "green": "548235",
    "gold": "B7791F",
    "orange": "C05621",
    "red": "B91C1C",
    "surface": "F8FAFC",
    "header": "D9EAF7",
    "subheader": "E2F0D9",
    "warn": "FCE4D6",
    "line": "CBD5E1",
}


def read(name):
    return pd.read_csv(AUDIT / name)


def safe_float(value, default=0.0):
    try:
        if pd.isna(value):
            return default
        return float(value)
    except Exception:
        return default


def esc(value):
    return html.escape(str(value), quote=True)


def fmt(value, digits=3):
    if pd.isna(value):
        return "NA"
    return f"{float(value):.{digits}f}"


def short_label(label):
    return str(label).replace("L0_", "L0 ").replace("L1_", "L1 ").replace("L2_", "L2 ").replace("L3_", "L3 ").replace("L4_", "L4 ").replace("L5_", "L5 ").replace("_", " ")


def write_df(ws, df, start_row=1, start_col=1, table_name=None):
    for j, col in enumerate(df.columns, start=start_col):
        cell = ws.cell(start_row, j, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=COLORS["blue"])
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False, name=None), start=start_row + 1):
        for j, value in enumerate(row, start=start_col):
            if isinstance(value, (np.integer,)):
                value = int(value)
            elif isinstance(value, (np.floating,)):
                value = float(value) if np.isfinite(value) else None
            elif pd.isna(value):
                value = None
            cell = ws.cell(i, j, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, float):
                cell.number_format = "0.000"
    end_row = start_row + len(df)
    end_col = start_col + len(df.columns) - 1
    thin = Side(style="thin", color=COLORS["line"])
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.border = Border(bottom=thin, right=thin)
    ws.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    return end_row, end_col


def set_widths(ws, df, max_width=42):
    for idx, col in enumerate(df.columns, start=1):
        series = df[col].astype(str).replace("nan", "")
        width = min(max(len(str(col)) + 2, int(series.str.len().quantile(0.90)) + 2 if len(series) else 12), max_width)
        if col.lower().endswith("id") or col in {"Article_ID", "GSE", "GPL"}:
            width = min(width, 20)
        if "interpretation" in col.lower() or "summary" in col.lower() or "caveat" in col.lower():
            width = max(width, 34)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def style_sheet(ws):
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Arial"
            font.sz = font.sz or 10
            cell.font = font


def add_title(ws, title, subtitle=None):
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=COLORS["blue"])
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    if subtitle:
        ws.merge_cells("A2:H2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color=COLORS["ink"])
        ws["A2"].fill = PatternFill("solid", fgColor=COLORS["surface"])
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 34


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    add_title(
        ws,
        "LUAD Precision-Oncology Deployability Audit",
        "This is not another LUAD prognostic model. It is a deployability audit of published transcriptomic prognostic signatures.",
    )
    readme_rows = [
        ("Scope", "178 full-text articles; 57 formula-reproducible signatures; 27 manuscript-triage signatures; 10 primary candidates; 9 GEO cohorts; TCGA GDC STAR TPM formula-direction sanity check."),
        ("Core question", "Do published LUAD transcriptomic prognostic signatures support individualized prognostic stratification, and where do formula reproducibility, score portability, cutoff transportability, and clinical deployability diverge?"),
        ("Primary strict-cutoff candidate", "A131 only."),
        ("Strong score-portability candidates", "A129, A067, A097, A168, A001, A059."),
        ("Terminology guardrail", "Median-fallback grouped GEO analyses are score-portability analyses, not original reported-cutoff validation."),
        ("TCGA guardrail", "TCGA-LUAD is a formula-direction sanity check, not a separate external-validation cohort."),
        ("Calibration guardrail", "Calibration metrics are external recalibration assessments; original absolute-risk calibration is not recoverable without published baseline survival functions."),
    ]
    ws.append([])
    ws.append(["Topic", "Positioning"])
    for r in readme_rows:
        ws.append(list(r))
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=COLORS["teal"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 118
    for row in ws.iter_rows(min_row=5, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    style_sheet(ws)

    sheets = {
        "Evidence_Ladder": read("evidence_ladder_summary_counts.csv"),
        "Reporting_Quality": read("reporting_quality_summary.csv"),
        "Cutoff_Audit": read("cutoff_transportability_summary.csv"),
        "Primary10_Checklist": read("clinical_deployability_checklist_primary10.csv"),
        "All57_Landscape": read("all_57_GEO_performance_summary.csv"),
        "Prediction_AUC": read("time_dependent_auc_primary_candidates.csv"),
        "Prediction_Brier": read("brier_score_primary_candidates.csv"),
        "Dataset_Reuse": read("dataset_reuse_frequency.csv").sort_values(["Signature_count", "Article_count"], ascending=False),
        "Platform": read("platform_stratified_performance.csv"),
        "Endpoint": read("endpoint_distribution_all_articles.csv"),
    }

    for name, df in sheets.items():
        ws = wb.create_sheet(name)
        if name == "Prediction_AUC":
            df = (
                df.groupby(["Article_ID", "Signature_ID", "time_years"], dropna=False)
                .agg(median_AUC=("dynamic_AUC_complete_case", "median"), evaluable_tests=("dynamic_AUC_complete_case", lambda s: s.notna().sum()))
                .reset_index()
            )
        elif name == "Prediction_Brier":
            df = (
                df.groupby(["Article_ID", "Signature_ID", "time_years"], dropna=False)
                .agg(median_Brier=("IPCW_Brier_score_external_recalibrated", "median"), evaluable_tests=("IPCW_Brier_score_external_recalibrated", lambda s: s.notna().sum()))
                .reset_index()
            )
        elif name == "Platform":
            df = (
                df.groupby(["Platform_class"], dropna=False)
                .agg(tests=("tests", "sum"), median_C_index=("median_C_index", "median"), median_group_HR=("median_group_HR", "median"), continuous_p05_rate=("continuous_p05_rate", "mean"))
                .reset_index()
            )
        elif name == "All57_Landscape":
            df = df.sort_values(["Evidence_ladder_level", "median_C_index"], ascending=[False, False])
        elif name == "Primary10_Checklist":
            order = {a: i for i, a in enumerate(PRIMARY_ARTICLE_ORDER)}
            df["_order"] = df["Article_ID"].map(order).fillna(99)
            df = df.sort_values("_order").drop(columns="_order")
        write_df(ws, df)
        set_widths(ws, df)
        style_sheet(ws)
        if name == "Primary10_Checklist":
            status_cols = [i + 1 for i, c in enumerate(df.columns) if c not in {"Article_ID", "Signature_ID", "Evidence_ladder_label", "Deployability_score_0_18", "Deployability_interpretation"}]
            fills = {
                "yes": PatternFill("solid", fgColor="D9EAD3"),
                "partial": PatternFill("solid", fgColor="FFF2CC"),
                "manual_check_required": PatternFill("solid", fgColor="D9EAF7"),
                "no": PatternFill("solid", fgColor="F4CCCC"),
            }
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for idx in status_cols:
                    value = str(row[idx - 1].value or "").lower()
                    if value in fills:
                        row[idx - 1].fill = fills[value]
        if name == "All57_Landscape":
            ci_col = list(df.columns).index("median_C_index") + 1
            ws.conditional_formatting.add(
                f"{get_column_letter(ci_col)}2:{get_column_letter(ci_col)}{ws.max_row}",
                ColorScaleRule(start_type="min", start_color="F4CCCC", mid_type="percentile", mid_value=50, mid_color="FFF2CC", end_type="max", end_color="D9EAD3"),
            )

    fig_index = wb.create_sheet("Figure_Index")
    fig_df = pd.DataFrame(
        [
            ("Figure 1", "Evidence ladder workflow", "figures/Figure1_evidence_ladder_workflow.svg"),
            ("Figure 2", "Reporting and reproducibility landscape", "figures/Figure2_reporting_reproducibility_landscape.svg"),
            ("Figure 3", "All-57 score-portability landscape", "figures/Figure3_all57_score_portability_landscape.svg"),
            ("Figure 4", "Cutoff transportability failure anatomy", "figures/Figure4_cutoff_transportability_failure_anatomy.svg"),
            ("Figure 5", "Primary candidate prediction performance", "figures/Figure5_primary_prediction_performance.svg"),
            ("Figure 6", "Clinical deployability gap heatmap", "figures/Figure6_clinical_deployability_gap_heatmap.svg"),
        ],
        columns=["Figure", "Purpose", "Relative_path"],
    )
    write_df(fig_index, fig_df)
    set_widths(fig_index, fig_df)
    style_sheet(fig_index)

    # Lightweight native charts for quick workbook scanning.
    ws = wb["Evidence_Ladder"]
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Evidence ladder counts"
    chart.y_axis.title = "Level"
    chart.x_axis.title = "Count"
    data = Reference(ws, min_col=3, min_row=1, max_row=ws.max_row)
    cats = Reference(ws, min_col=2, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "E3")

    ws = wb["Prediction_AUC"]
    chart = BarChart()
    chart.title = "Median time-dependent AUC by signature/time"
    chart.y_axis.title = "Median AUC"
    chart.height = 8
    chart.width = 16
    # Chart the first compact block only to avoid visual overload.
    max_row = min(ws.max_row, 31)
    data = Reference(ws, min_col=4, min_row=1, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "G3")

    wb.save(XLSX)
    # Verify workbook is readable.
    check = load_workbook(XLSX, read_only=True, data_only=False)
    assert "README" in check.sheetnames and "Primary10_Checklist" in check.sheetnames
    check.close()


def svg_text(x, y, text, size=13, fill=None, weight="normal", anchor="start"):
    return f'<text x="{x}" y="{y}" font-family="Arial" font-size="{size}" font-weight="{weight}" text-anchor="{anchor}" fill="#{fill or COLORS["ink"]}">{esc(text)}</text>'


def svg_rect(x, y, w, h, fill, stroke=None, rx=4, opacity=1.0):
    stroke_attr = f' stroke="#{stroke}"' if stroke else ""
    return f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{rx}" fill="#{fill}" opacity="{opacity}"{stroke_attr}/>'


def save_svg(name, body, width=1100, height=720):
    svg = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        svg_rect(0, 0, width, height, "FFFFFF", rx=0),
        body,
        "</svg>",
    ]
    path = FIG_DIR / name
    path.write_text("\n".join(svg), encoding="utf-8")


def build_figures():
    FIG_DIR.mkdir(parents=True, exist_ok=True)

    # Figure 1: ladder workflow.
    counts = read("evidence_ladder_summary_counts.csv")
    sig = counts[counts["Entity"].eq("signature")].copy()
    labels = list(sig["Evidence_ladder_label"])
    values = list(sig["Count"])
    x0, y0, box_w, box_h, gap = 35, 170, 155, 100, 20
    body = [svg_text(35, 50, "Figure 1. Evidence ladder workflow", 22, COLORS["ink"], "bold")]
    body.append(svg_text(35, 80, "Attrition from formula reproducibility to deployability readiness", 13, COLORS["muted"]))
    for i, (label, value) in enumerate(zip(labels, values)):
        x = x0 + i * (box_w + gap)
        fill = ["F4CCCC", "D9EAF7", "D9EAF7", "D9EAD3", "FFF2CC", "EADCF8"][i]
        body.append(svg_rect(x, y0, box_w, box_h, fill, COLORS["line"], rx=8))
        wrapped = textwrap.wrap(short_label(label), width=19)
        for k, line in enumerate(wrapped[:3]):
            body.append(svg_text(x + box_w / 2, y0 + 28 + k * 16, line, 12, COLORS["ink"], "bold" if k == 0 else "normal", "middle"))
        body.append(svg_text(x + box_w / 2, y0 + 82, str(int(value)), 24, COLORS["blue"], "bold", "middle"))
        if i < len(labels) - 1:
            ax = x + box_w + 4
            ay = y0 + box_h / 2
            body.append(f'<line x1="{ax}" y1="{ay}" x2="{ax + gap - 8}" y2="{ay}" stroke="#{COLORS["muted"]}" stroke-width="2"/>')
            body.append(f'<polygon points="{ax + gap - 8},{ay - 5} {ax + gap},{ay} {ax + gap - 8},{ay + 5}" fill="#{COLORS["muted"]}"/>')
    body.append(svg_text(35, 335, "Interpretation: evidence levels are not interchangeable; score portability does not imply cutoff transportability.", 14, COLORS["ink"], "bold"))
    save_svg("Figure1_evidence_ladder_workflow.svg", "\n".join(body), 1120, 420)

    # Figure 2: reporting/reproducibility.
    year = read("reporting_score_by_year.csv").dropna(subset=["Year"]).sort_values("Year")
    years = year["Year"].astype(int).tolist()
    scores = year["Mean_reporting_score"].tolist()
    repro = year["Formula_reproducible_count"].tolist()
    w, h = 1050, 620
    px0, py0, pw, ph = 80, 100, 850, 360
    body = [svg_text(40, 48, "Figure 2. Reporting and reproducibility landscape", 22, COLORS["ink"], "bold")]
    body.append(svg_text(40, 76, "Reporting quality improved unevenly; formula reproducibility remains a distinct hurdle.", 13, COLORS["muted"]))
    body.append(f'<line x1="{px0}" y1="{py0+ph}" x2="{px0+pw}" y2="{py0+ph}" stroke="#{COLORS["ink"]}"/>')
    body.append(f'<line x1="{px0}" y1="{py0}" x2="{px0}" y2="{py0+ph}" stroke="#{COLORS["ink"]}"/>')
    max_score = 10
    for i, (yr, sc, rp) in enumerate(zip(years, scores, repro)):
        x = px0 + i * (pw / max(1, len(years) - 1))
        y = py0 + ph - (sc / max_score) * ph
        r = 5 + min(rp, 8)
        body.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{r}" fill="#{COLORS["teal"]}" opacity="0.8"/>')
        body.append(svg_text(x, py0 + ph + 22, str(yr), 10, COLORS["muted"], anchor="middle"))
    points = []
    for i, sc in enumerate(scores):
        x = px0 + i * (pw / max(1, len(scores) - 1))
        y = py0 + ph - (sc / max_score) * ph
        points.append(f"{x:.1f},{y:.1f}")
    body.append(f'<polyline points="{" ".join(points)}" fill="none" stroke="#{COLORS["blue"]}" stroke-width="2"/>')
    for tick in range(0, 11, 2):
        y = py0 + ph - tick / 10 * ph
        body.append(svg_text(px0 - 12, y + 4, str(tick), 10, COLORS["muted"], anchor="end"))
        body.append(f'<line x1="{px0-4}" y1="{y}" x2="{px0+pw}" y2="{y}" stroke="#EEF2F7"/>')
    body.append(svg_text(40, 515, "Dot size encodes formula-reproducible article count. Line shows mean 0-10 reporting score.", 12, COLORS["muted"]))
    save_svg("Figure2_reporting_reproducibility_landscape.svg", "\n".join(body), w, h)

    # Figure 3: all-57 landscape.
    perf = read("all_57_GEO_performance_summary.csv")
    primary = set(PRIMARY_ARTICLE_ORDER)
    body = [svg_text(40, 48, "Figure 3. All-57 external score-portability landscape", 22, COLORS["ink"], "bold")]
    body.append(svg_text(40, 76, "External signal is a minority phenotype among formula-reproducible signatures.", 13, COLORS["muted"]))
    px0, py0, pw, ph = 95, 110, 820, 390
    body.append(f'<line x1="{px0}" y1="{py0+ph}" x2="{px0+pw}" y2="{py0+ph}" stroke="#{COLORS["ink"]}"/>')
    body.append(f'<line x1="{px0}" y1="{py0}" x2="{px0}" y2="{py0+ph}" stroke="#{COLORS["ink"]}"/>')
    for _, r in perf.iterrows():
        x = px0 + np.clip((safe_float(r["median_C_index"]) - 0.45) / 0.30, 0, 1) * pw
        y = py0 + ph - np.clip(safe_float(r["continuous_significant_rate"]), 0, 1) * ph
        is_primary = r["Article_ID"] in primary
        color = COLORS["orange"] if is_primary else COLORS["teal"]
        radius = 8 if is_primary else 5
        body.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{radius}" fill="#{color}" opacity="{0.88 if is_primary else 0.45}"/>')
        if is_primary:
            body.append(svg_text(x + 8, y - 7, r["Article_ID"], 10, COLORS["ink"], "bold"))
    for val in [0.5, 0.55, 0.60, 0.65, 0.70]:
        x = px0 + (val - 0.45) / 0.30 * pw
        body.append(svg_text(x, py0 + ph + 24, f"{val:.2f}", 10, COLORS["muted"], anchor="middle"))
    for val in [0, 0.25, 0.5, 0.75, 1.0]:
        y = py0 + ph - val * ph
        body.append(svg_text(px0 - 12, y + 4, f"{val:.2f}", 10, COLORS["muted"], anchor="end"))
        body.append(f'<line x1="{px0}" y1="{y}" x2="{px0+pw}" y2="{y}" stroke="#EEF2F7"/>')
    body.append(svg_text(px0 + pw / 2, py0 + ph + 55, "Median C-index across evaluable GEO cohorts", 13, COLORS["ink"], anchor="middle"))
    body.append(svg_text(25, py0 + 170, "Continuous p<0.05 rate", 13, COLORS["ink"]))
    save_svg("Figure3_all57_score_portability_landscape.svg", "\n".join(body), 1040, 620)

    # Figure 4: cutoff anatomy.
    cutoff = read("cutoff_transportability_summary.csv")
    body = [svg_text(40, 48, "Figure 4. Cutoff transportability failure anatomy", 22, COLORS["ink"], "bold")]
    body.append(svg_text(40, 76, "Reported numeric thresholds are uncommon and often do not survive cross-platform scale transfer.", 13, COLORS["muted"]))
    labels = []
    vals = []
    for _, r in cutoff.iterrows():
        label = (
            f"numeric={r['reported_numeric_cutoff']}; applicable={r['reported_cutoff_applicable_to_GEO']}; collapse={r['single_group_collapse_observed']}"
        )
        labels.append(label)
        vals.append(int(r["Signature_count"]))
    x0, y0, maxw, bh, gap = 360, 125, 540, 42, 18
    maxv = max(vals) if vals else 1
    for i, (label, val) in enumerate(zip(labels, vals)):
        y = y0 + i * (bh + gap)
        body.append(svg_text(42, y + 26, label, 12, COLORS["ink"]))
        bw = maxw * val / maxv
        fill = COLORS["green"] if "applicable=True; collapse=False" in label else COLORS["orange"] if "numeric=True" in label else COLORS["muted"]
        body.append(svg_rect(x0, y, bw, bh, fill, rx=4, opacity=0.82))
        body.append(svg_text(x0 + bw + 10, y + 27, str(val), 13, COLORS["ink"], "bold"))
    body.append(svg_text(42, 470, "A097 manual PDF cutoff 1.66 is counted as reported but non-transportable: fixed cutoff collapses all tested GEO cohorts.", 13, COLORS["red"], "bold"))
    save_svg("Figure4_cutoff_transportability_failure_anatomy.svg", "\n".join(body), 1040, 560)

    # Figure 5: AUC/Brier by primary candidate.
    auc = read("time_dependent_auc_primary_candidates.csv")
    brier = read("brier_score_primary_candidates.csv")
    auc3 = auc[auc["time_years"].eq(3.0)].groupby("Article_ID")["dynamic_AUC_complete_case"].median()
    auc5 = auc[auc["time_years"].eq(5.0)].groupby("Article_ID")["dynamic_AUC_complete_case"].median()
    br3 = brier[brier["time_years"].eq(3.0)].groupby("Article_ID")["IPCW_Brier_score_external_recalibrated"].median()
    articles = [a for a in PRIMARY_ARTICLE_ORDER if a in auc3.index]
    body = [svg_text(40, 48, "Figure 5. Primary candidate prediction performance", 22, COLORS["ink"], "bold")]
    body.append(svg_text(40, 76, "Time-dependent AUC and external recalibration Brier scores summarize prediction behavior, not deployment readiness.", 13, COLORS["muted"]))
    x0, y0, pw, ph = 75, 120, 880, 340
    bar_w = pw / len(articles) / 3
    for i, a in enumerate(articles):
        base = x0 + i * (pw / len(articles)) + 10
        h3 = (safe_float(auc3.get(a), 0.5) - 0.45) / 0.35 * ph
        h5 = (safe_float(auc5.get(a), 0.5) - 0.45) / 0.35 * ph
        body.append(svg_rect(base, y0 + ph - h3, bar_w, h3, COLORS["blue"], rx=2))
        body.append(svg_rect(base + bar_w + 3, y0 + ph - h5, bar_w, h5, COLORS["teal"], rx=2))
        body.append(svg_text(base + bar_w, y0 + ph + 20, a, 10, COLORS["ink"], "bold", "middle"))
        body.append(svg_text(base + bar_w, y0 + ph + 36, f"Brier3 {fmt(br3.get(a), 2)}", 9, COLORS["muted"], anchor="middle"))
    for val in [0.5, 0.6, 0.7, 0.8]:
        y = y0 + ph - (val - 0.45) / 0.35 * ph
        body.append(f'<line x1="{x0}" y1="{y}" x2="{x0+pw}" y2="{y}" stroke="#EEF2F7"/>')
        body.append(svg_text(x0 - 10, y + 4, f"{val:.1f}", 10, COLORS["muted"], anchor="end"))
    body.append(svg_text(760, 105, "Blue: 3-year AUC   Teal: 5-year AUC", 12, COLORS["ink"]))
    save_svg("Figure5_primary_prediction_performance.svg", "\n".join(body), 1040, 560)

    # Figure 6: deployability heatmap.
    heat = read("clinical_deployability_heatmap_data.csv")
    domains = list(heat["Deployability_domain"].drop_duplicates())
    articles = [a for a in PRIMARY_ARTICLE_ORDER if a in set(heat["Article_ID"])]
    cell_w, cell_h = 36, 28
    left, top = 170, 130
    body = [svg_text(35, 48, "Figure 6. Clinical deployability gap heatmap", 22, COLORS["ink"], "bold")]
    body.append(svg_text(35, 76, "Most candidates retain prognostic signal but lack deployable clinical-threshold infrastructure.", 13, COLORS["muted"]))
    val_map = {(r.Article_ID, r.Deployability_domain): r.Heatmap_value for r in heat.itertuples()}
    for i, dom in enumerate(domains):
        x = left + i * cell_w
        body.append(f'<text x="{x+12}" y="{top-12}" font-family="Arial" font-size="9" fill="#{COLORS["muted"]}" transform="rotate(-55 {x+12},{top-12})">{esc(dom)}</text>')
    for j, article in enumerate(articles):
        y = top + j * cell_h
        body.append(svg_text(left - 12, y + 18, article, 11, COLORS["ink"], "bold", "end"))
        for i, dom in enumerate(domains):
            value = safe_float(val_map.get((article, dom), 0))
            fill = "D9EAD3" if value >= 0.99 else "FFF2CC" if value >= 0.49 else "D9EAF7" if value >= 0.24 else "F4CCCC"
            body.append(svg_rect(left + i * cell_w, y, cell_w - 2, cell_h - 2, fill, "FFFFFF", rx=1))
    body.append(svg_text(35, 520, "Green=yes; yellow=partial; blue=manual check; red=no.", 12, COLORS["muted"]))
    save_svg("Figure6_clinical_deployability_gap_heatmap.svg", "\n".join(body), 930, 560)


def build_manuscript():
    ladder = read("evidence_ladder_summary_counts.csv")
    sig_counts = ladder[ladder["Entity"].eq("signature")].set_index("Evidence_ladder_label")["Count"].to_dict()
    deploy = read("clinical_deployability_score_summary.csv")
    cutoff = read("cutoff_transportability_summary.csv")
    auc = read("time_dependent_auc_primary_candidates.csv")
    brier = read("brier_score_primary_candidates.csv")
    auc_medians = auc.groupby("time_years")["dynamic_AUC_complete_case"].median().to_dict()
    brier_medians = brier.groupby("time_years")["IPCW_Brier_score_external_recalibrated"].median().to_dict()

    text = f"""# A precision-oncology deployability audit of published LUAD transcriptomic prognostic signatures

## Positioning

This is not another LUAD prognostic model. It is a precision-oncology deployability audit of published transcriptomic prognostic signatures.

The central question is whether published LUAD transcriptomic prognostic signatures truly support individualized prognostic stratification, and where formula reproducibility, external scoreability, score portability, cutoff transportability, and clinical deployability diverge.

## Abstract Draft

Published transcriptomic prognostic signatures are frequently presented as tools for risk stratification in lung adenocarcinoma (LUAD), but the evidence required for precision-oncology use extends beyond nominal prognostic association. We audited 178 full-text LUAD prognostic-signature articles, identifying 57 formula-reproducible signatures and evaluating their external scoreability and transportability across nine GEO cohorts, with TCGA-LUAD GDC STAR TPM rescoring used only as a formula-direction sanity check. Evidence was organized into a ladder from non-reproducible reporting to clinical deployability candidacy.

Of 57 formula-reproducible signatures, {int(sig_counts.get('L3_score_portable', 0))} reached score-portability level, whereas only A131 emerged as the strict reported-cutoff primary candidate. A129, A067, A097, A168, A001, and A059 showed strong score-portability evidence, but their grouped GEO analyses relied on cohort-median fallback and therefore do not test original reported thresholds. A097 illustrates the cutoff bottleneck: manual PDF verification confirmed a TCGA/training cutoff of 1.66, but direct application to processed GEO matrices collapsed all tested cohorts into a single risk group. Across primary candidates, median time-dependent AUC was {fmt(auc_medians.get(3.0))} at 3 years and {fmt(auc_medians.get(5.0))} at 5 years; median external-recalibration IPCW Brier score was {fmt(brier_medians.get(3.0))} at 3 years and {fmt(brier_medians.get(5.0))} at 5 years.

These findings show that prognostic signal, cutoff transportability, and deployability are separable evidence levels. Published LUAD transcriptomic signatures may retain portable biology, but most do not yet provide the fixed thresholds, assay standardization, prospective evidence, or absolute-risk calibration information required for precision-oncology deployment.

## Results

### Published LUAD prognostic signatures frequently lack complete reproducibility elements

The audit began with 178 full-text articles and separated article-level reporting from signature-level formula reproducibility. The evidence ladder identified 57 formula-reproducible signatures, while 121 articles remained at L0 because one or more essential reproducibility elements were insufficiently reported. The reporting quality score captures gene list, coefficients, formula, numeric cutoff, cutoff derivation, training cohort, validation cohort, endpoint, preprocessing, and code/data availability.

### External scoreability is limited by gene coverage, cohort reuse, and endpoint inconsistency

All formula-reproducible signatures were tested against the same GEO validation framework when gene coverage and cohort-overlap rules permitted. Failures were not treated as biological invalidation; they were evidence-level failures arising from missing genes on external platforms, training overlap, single-group collapse under fixed thresholds, or incomplete endpoint alignment with GEO OS.

### Score portability is preserved in a minority of signatures

Score portability was strongest for A129, A067, A097, A168, A001, and A059, with A098 showing moderate support and A065/A173 lower-priority support. These claims are score-portability claims. Median-fallback high/low grouping is a cohort-specific analysis convenience and should not be written as reported-cutoff validation.

### Reported cutoff transportability is rare and represents the major clinical translation bottleneck

Cutoff reporting and cutoff transportability were audited separately. A131 is the only primary strict reported-cutoff candidate. Numeric thresholds were often missing, insufficiently tied to expression scale, or non-transportable across GEO preprocessing contexts. A097 is explicitly handled as a reported-cutoff failure case: the published TCGA/training cutoff exists, but direct GEO application produced single-group collapse.

### Clinical adjustment, expression-scale sensitivity, and TCGA sanity checks support selected signals but do not establish deployability

Clinical-adjusted Cox models, clinical-only versus clinical-plus-risk C-index comparisons, gene-zscore sensitivity, leave-one-cohort-out checks, and TCGA GDC STAR TPM rescoring support selected formula directions. TCGA-LUAD is interpreted as a formula-direction sanity check because many signatures used TCGA during development.

### Evidence ladder and deployability checklist reveal a gap between prognostic biology and precision-oncology readiness

The clinical deployability checklist shows that A131 is the best candidate for further deployability evaluation, but not ready for deployment. Common remaining gaps include original code/calculator availability, assay standardization, prospective validation, clinical threshold justification, and original absolute-risk calibration. Calibration outputs in this package are external recalibration assessments only.

## Figure Plan

1. Figure 1: Evidence ladder workflow.
2. Figure 2: Reporting and reproducibility landscape.
3. Figure 3: All-57 external score-portability landscape.
4. Figure 4: Cutoff transportability failure anatomy.
5. Figure 5: Primary candidate prediction performance.
6. Figure 6: Clinical deployability gap heatmap.

## Terminology Rules for Final Drafting

- Use reported-cutoff validation only when the original reported numeric cutoff is applied.
- Use score-portability analysis for median-fallback grouped GEO analyses.
- Describe TCGA-LUAD as a formula-direction sanity check.
- Describe pooled HR as descriptive pooled HR unless a full formal meta-analysis method is added.
- Avoid valid/invalid binary framing; prognostic signal, cutoff transportability, and clinical deployability are separable evidence levels.
"""
    MANUSCRIPT.write_text(text, encoding="utf-8")


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color="1F2937"):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_df_table(doc, df, columns, widths, max_rows=None):
    view = df[columns].head(max_rows).copy() if max_rows else df[columns].copy()
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    for i, col in enumerate(columns):
        table.columns[i].width = Inches(widths[i])
        set_cell_text(table.rows[0].cells[i], col, bold=True, color="FFFFFF")
        set_cell_fill(table.rows[0].cells[i], COLORS["blue"])
    for row in view.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if isinstance(value, float):
                value = fmt(value)
            set_cell_text(cells[i], "" if pd.isna(value) else value)
    return table


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor.from_string(COLORS["blue"])
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = RGBColor.from_string(COLORS["teal"])
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("A precision-oncology deployability audit of published LUAD transcriptomic prognostic signatures")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLORS["ink"])

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Working manuscript draft | generated from the deployability audit package")
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    callout = doc.add_paragraph()
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(10)
    run = callout.add_run("Positioning: ")
    run.bold = True
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(COLORS["blue"])
    run = callout.add_run("This is not another LUAD prognostic model. It is a precision-oncology deployability audit of published transcriptomic prognostic signatures.")
    run.font.name = "Arial"
    run.font.size = Pt(10.5)

    md = MANUSCRIPT.read_text(encoding="utf-8")
    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\. ", "", line))
        else:
            doc.add_paragraph(line)

    doc.add_page_break()
    doc.add_heading("Primary Candidate Deployability Summary", level=1)
    deploy = read("clinical_deployability_score_summary.csv")
    order = {a: i for i, a in enumerate(PRIMARY_ARTICLE_ORDER)}
    deploy["_order"] = deploy["Article_ID"].map(order).fillna(99)
    deploy = deploy.sort_values("_order")
    add_df_table(
        doc,
        deploy,
        ["Article_ID", "Signature_ID", "Deployability_score_0_18", "Deployability_interpretation"],
        [0.75, 1.65, 0.9, 3.05],
        max_rows=10,
    )

    doc.add_heading("Recommended Figure Set", level=1)
    fig_df = pd.DataFrame(
        [
            ("Figure 1", "Evidence ladder workflow"),
            ("Figure 2", "Reporting and reproducibility landscape"),
            ("Figure 3", "All-57 external score-portability landscape"),
            ("Figure 4", "Cutoff transportability failure anatomy"),
            ("Figure 5", "Primary candidate prediction performance"),
            ("Figure 6", "Clinical deployability gap heatmap"),
        ],
        columns=["Figure", "Purpose"],
    )
    add_df_table(doc, fig_df, ["Figure", "Purpose"], [1.0, 5.2])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("LUAD deployability audit draft")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    doc.save(DOCX)
    check = Document(DOCX)
    assert len(check.paragraphs) > 20


def build_package():
    if PKG_DIR.exists():
        shutil.rmtree(PKG_DIR)
    (PKG_DIR / "tables").mkdir(parents=True)
    (PKG_DIR / "figures").mkdir(parents=True)
    (PKG_DIR / "notes").mkdir(parents=True)
    (PKG_DIR / "data").mkdir(parents=True)

    shutil.copy2(XLSX, PKG_DIR / "tables" / XLSX.name)
    shutil.copy2(MANUSCRIPT, PKG_DIR / "notes" / MANUSCRIPT.name)
    shutil.copy2(DOCX, PKG_DIR / "notes" / DOCX.name)
    for md in [
        "npj_precision_oncology_readiness_report.md",
        "high_impact_analysis_summary.md",
        "clinical_deployability_gap_summary.md",
        "dataset_reuse_summary.md",
        "endpoint_consistency_summary.md",
        "npj_precision_oncology_manuscript_outline.md",
    ]:
        src = AUDIT / md
        if src.exists():
            shutil.copy2(src, PKG_DIR / "notes" / md)
    for fig in FIG_DIR.glob("*.svg"):
        shutil.copy2(fig, PKG_DIR / "figures" / fig.name)
    for csv in AUDIT.glob("*.csv"):
        # Keep full CSV data pack for reproducibility.
        shutil.copy2(csv, PKG_DIR / "data" / csv.name)
    if ZIP.exists():
        ZIP.unlink()
    with zipfile.ZipFile(ZIP, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in PKG_DIR.rglob("*"):
            if path.is_file():
                zf.write(path, path.relative_to(PKG_DIR.parent))


def validate_outputs():
    required = [
        XLSX,
        MANUSCRIPT,
        DOCX,
        ZIP,
        FIG_DIR / "Figure1_evidence_ladder_workflow.svg",
        FIG_DIR / "Figure2_reporting_reproducibility_landscape.svg",
        FIG_DIR / "Figure3_all57_score_portability_landscape.svg",
        FIG_DIR / "Figure4_cutoff_transportability_failure_anatomy.svg",
        FIG_DIR / "Figure5_primary_prediction_performance.svg",
        FIG_DIR / "Figure6_clinical_deployability_gap_heatmap.svg",
    ]
    missing = [str(p) for p in required if not p.exists() or p.stat().st_size == 0]
    if missing:
        raise RuntimeError(f"Missing outputs: {missing}")
    text = MANUSCRIPT.read_text(encoding="utf-8").lower()
    banned = ["independent external validation", "published signatures are invalid", "ordinary bioinformatics", "deployment-ready"]
    hits = [b for b in banned if b in text]
    if hits:
        raise RuntimeError(f"Terminology guardrail failed: {hits}")
    wb = load_workbook(XLSX, read_only=True, data_only=False)
    expected_sheets = {"README", "Evidence_Ladder", "Primary10_Checklist", "All57_Landscape", "Prediction_AUC", "Figure_Index"}
    missing_sheets = expected_sheets - set(wb.sheetnames)
    wb.close()
    if missing_sheets:
        raise RuntimeError(f"Workbook missing sheets: {missing_sheets}")
    doc = Document(DOCX)
    if len(doc.paragraphs) < 20:
        raise RuntimeError("DOCX structural check failed: too few paragraphs")


def main():
    AUDIT.mkdir(parents=True, exist_ok=True)
    build_workbook()
    build_figures()
    build_manuscript()
    build_docx()
    build_package()
    validate_outputs()
    print(f"Workbook: {XLSX}")
    print(f"Manuscript: {MANUSCRIPT}")
    print(f"Figures: {FIG_DIR}")
    print(f"Package: {ZIP}")


if __name__ == "__main__":
    main()
