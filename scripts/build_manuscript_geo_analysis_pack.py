import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
OUT_DIR = ROOT / "outputs" / "luad_signature_extraction"

SELECTION = WORK / "luad_geo_manuscript_ready_signature_selection.csv"
RESULTS = WORK / "geo_external_validation_results.csv"
RESULTS_OVERRIDE = WORK / "geo_external_validation_results.manual_overrides.csv"

OUT_SIGNATURE = WORK / "luad_geo_manuscript_signature_evidence.csv"
OUT_PAIRWISE = WORK / "luad_geo_manuscript_pairwise_evidence.csv"
OUT_MATRIX = WORK / "luad_geo_manuscript_cohort_matrix.csv"
OUT_META = WORK / "luad_geo_manuscript_pooled_hr.csv"
OUT_DECISION = WORK / "luad_geo_manuscript_decision_ranking.csv"
OUT_MD = OUT_DIR / "LUAD_GEO_manuscript_results_draft.md"
OUT_XLSX = OUT_DIR / "LUAD_GEO_manuscript_analysis_pack.xlsx"
OUT_HEATMAP = OUT_DIR / "LUAD_GEO_main_signature_heatmap.svg"
OUT_FOREST = OUT_DIR / "LUAD_GEO_main_signature_pooled_HR_forest.svg"


MAIN_TIERS = {
    "Tier_1_strict_reported_cutoff",
    "Tier_2_primary_median_fallback",
}

SENSITIVITY_TIERS = {
    "Tier_3_secondary_median_fallback",
    "Tier_4_continuous_score_support",
    "Tier_5_manual_override_sensitivity",
}

TIER_ORDER = {
    "Tier_1_strict_reported_cutoff": 1,
    "Tier_2_primary_median_fallback": 2,
    "Tier_3_secondary_median_fallback": 3,
    "Tier_4_continuous_score_support": 4,
    "Tier_5_manual_override_sensitivity": 5,
}

GSE_ORDER = [
    "GSE13213",
    "GSE30219",
    "GSE31210",
    "GSE3141",
    "GSE37745",
    "GSE42127",
    "GSE50081",
    "GSE68465",
    "GSE72094",
]


def num(series):
    return pd.to_numeric(series, errors="coerce")


def normal_p_two_sided(z):
    if not math.isfinite(z):
        return np.nan
    return math.erfc(abs(z) / math.sqrt(2.0))


def fixed_effect_meta(rows):
    rows = rows.copy()
    rows["HR"] = num(rows["Cox_high_vs_low_HR"])
    rows["CI_low"] = num(rows["Cox_high_vs_low_CI95_low"])
    rows["CI_high"] = num(rows["Cox_high_vs_low_CI95_high"])
    rows = rows[(rows["HR"] > 0) & (rows["CI_low"] > 0) & (rows["CI_high"] > rows["CI_low"])]
    if len(rows) == 0:
        return {}
    y = np.log(rows["HR"].to_numpy(float))
    se = (np.log(rows["CI_high"].to_numpy(float)) - np.log(rows["CI_low"].to_numpy(float))) / (2 * 1.96)
    valid = np.isfinite(y) & np.isfinite(se) & (se > 0)
    y, se = y[valid], se[valid]
    if len(y) == 0:
        return {}
    w = 1.0 / (se**2)
    pooled = np.sum(w * y) / np.sum(w)
    pooled_se = math.sqrt(1.0 / np.sum(w))
    ci_low = pooled - 1.96 * pooled_se
    ci_high = pooled + 1.96 * pooled_se
    q = float(np.sum(w * (y - pooled) ** 2))
    df = len(y) - 1
    i2 = max(0.0, (q - df) / q) * 100 if q > 0 and df > 0 else 0.0
    return {
        "Pooled_HR_fixed": math.exp(pooled),
        "Pooled_HR_CI95_low": math.exp(ci_low),
        "Pooled_HR_CI95_high": math.exp(ci_high),
        "Pooled_HR_p": normal_p_two_sided(pooled / pooled_se),
        "Meta_k": int(len(y)),
        "Meta_I2_percent": i2,
    }


def bh_fdr(values):
    vals = pd.to_numeric(pd.Series(values), errors="coerce")
    out = pd.Series([np.nan] * len(vals), index=vals.index, dtype=float)
    valid = vals.dropna()
    if valid.empty:
        return out
    order = valid.sort_values().index
    m = len(valid)
    prev = 1.0
    for rank, idx in reversed(list(enumerate(order, start=1))):
        q = min(prev, valid.loc[idx] * m / rank)
        out.loc[idx] = min(q, 1.0)
        prev = q
    return out


def load_pairwise(selection):
    base = pd.read_csv(RESULTS)
    override = pd.read_csv(RESULTS_OVERRIDE)
    manual_sigs = {"A010_PMID40591077", "A104_PMID34760888"}
    base["Analysis_source"] = "first_pass"
    override = override[override["Signature_ID"].isin(manual_sigs)].copy()
    override["Analysis_source"] = "manual_override_sensitivity"
    combined = pd.concat([base[~base["Signature_ID"].isin(manual_sigs)], override], ignore_index=True)
    keep = set(selection["Signature_ID"])
    combined = combined[combined["Signature_ID"].isin(keep)].copy()
    combined = combined.merge(
        selection[
            [
                "Signature_ID",
                "Article_ID",
                "PMID",
                "Final_analysis_tier",
                "Recommended_manuscript_role",
                "Key_caveats_for_manuscript",
                "Original_validation_GSEs_tested",
            ]
        ],
        on=["Signature_ID", "Article_ID", "PMID"],
        how="left",
        suffixes=("", "_selection"),
    )
    for col in [
        "N",
        "Events",
        "Gene_coverage",
        "High_n",
        "Low_n",
        "C_index",
        "Cox_continuous_HR_per_SD",
        "Cox_continuous_p",
        "Cox_high_vs_low_HR",
        "Cox_high_vs_low_CI95_low",
        "Cox_high_vs_low_CI95_high",
        "Cox_high_vs_low_p",
        "Logrank_p",
    ]:
        if col in combined:
            combined[col] = num(combined[col])
    combined["Group_sig_p05"] = combined["Cox_high_vs_low_p"].lt(0.05)
    combined["Continuous_sig_p05"] = combined["Cox_continuous_p"].lt(0.05)
    combined["Group_HR_gt1"] = combined["Cox_high_vs_low_HR"].gt(1)
    combined["Continuous_HR_gt1"] = combined["Cox_continuous_HR_per_SD"].gt(1)
    combined["Is_main_analysis"] = combined["Final_analysis_tier"].isin(MAIN_TIERS)
    combined["Is_sensitivity_analysis"] = combined["Final_analysis_tier"].isin(SENSITIVITY_TIERS)
    combined["Uses_GSE50081"] = combined["GSE"].eq("GSE50081")
    combined["GSE_was_original_validation"] = combined.apply(
        lambda r: str(r["GSE"]) in str(r.get("Original_validation_GSEs_tested", "")).split("; "), axis=1
    )
    combined["Group_p_BH_selected_ok"] = np.nan
    combined["Continuous_p_BH_selected_ok"] = np.nan
    combined["Group_p_BH_main_ok"] = np.nan
    combined["Continuous_p_BH_main_ok"] = np.nan
    ok_idx = combined.index[combined["status"].eq("ok")]
    combined.loc[ok_idx, "Group_p_BH_selected_ok"] = bh_fdr(combined.loc[ok_idx, "Cox_high_vs_low_p"]).values
    combined.loc[ok_idx, "Continuous_p_BH_selected_ok"] = bh_fdr(combined.loc[ok_idx, "Cox_continuous_p"]).values
    main_idx = combined.index[combined["status"].eq("ok") & combined["Is_main_analysis"]]
    combined.loc[main_idx, "Group_p_BH_main_ok"] = bh_fdr(combined.loc[main_idx, "Cox_high_vs_low_p"]).values
    combined.loc[main_idx, "Continuous_p_BH_main_ok"] = bh_fdr(combined.loc[main_idx, "Cox_continuous_p"]).values
    return combined


def build_signature_evidence(selection, pairwise):
    ok = pairwise[pairwise["status"].eq("ok")].copy()
    rows = []
    for _, srow in selection.iterrows():
        sig = srow["Signature_ID"]
        sub = ok[ok["Signature_ID"].eq(sig)].copy()
        meta = fixed_effect_meta(sub)
        group_sig = int(sub["Group_sig_p05"].sum()) if len(sub) else 0
        cont_sig = int(sub["Continuous_sig_p05"].sum()) if len(sub) else 0
        events = int(num(sub["Events"]).sum()) if len(sub) else 0
        best = ""
        if len(sub) and sub["C_index"].notna().any():
            best_row = sub.loc[sub["C_index"].idxmax()]
            best = f"{best_row['GSE']} (C-index {best_row['C_index']:.3f})"
        rows.append(
            {
                "Final_analysis_tier": srow["Final_analysis_tier"],
                "Article_ID": srow["Article_ID"],
                "Signature_ID": sig,
                "PMID": srow["PMID"],
                "Main_or_sensitivity": "main" if srow["Final_analysis_tier"] in MAIN_TIERS else "sensitivity",
                "Cutoff_basis_for_GEO": "reported numeric cutoff"
                if srow["Final_analysis_tier"] == "Tier_1_strict_reported_cutoff"
                else ("manual override sensitivity" if srow["Final_analysis_tier"] == "Tier_5_manual_override_sensitivity" else "cohort median fallback / continuous score"),
                "GEO_tests_ok": int(len(sub)),
                "Total_events": events,
                "Group_sig_p05_n": group_sig,
                "Group_sig_BH_main_q05_n": int(sub["Group_p_BH_main_ok"].lt(0.05).sum()) if len(sub) else 0,
                "Group_sig_p05_rate": group_sig / len(sub) if len(sub) else np.nan,
                "Continuous_sig_p05_n": cont_sig,
                "Continuous_sig_BH_main_q05_n": int(sub["Continuous_p_BH_main_ok"].lt(0.05).sum()) if len(sub) else 0,
                "Continuous_sig_p05_rate": cont_sig / len(sub) if len(sub) else np.nan,
                "HR_direction_gt1_n": int(sub["Group_HR_gt1"].sum()) if len(sub) else 0,
                "HR_direction_gt1_rate": sub["Group_HR_gt1"].mean() if len(sub) else np.nan,
                "Median_group_HR": float(sub["Cox_high_vs_low_HR"].median()) if len(sub) else np.nan,
                "Median_C_index": float(sub["C_index"].median()) if len(sub) else np.nan,
                "Best_GEO_by_C_index": best,
                "GSEs_ok": "; ".join([g for g in GSE_ORDER if g in set(sub["GSE"])]),
                "GSE50081_included": "yes" if "GSE50081" in set(sub["GSE"]) else "no",
                "Original_validation_GSEs_retested": srow.get("Original_validation_GSEs_tested", ""),
                "Key_caveats_for_manuscript": srow.get("Key_caveats_for_manuscript", ""),
                **meta,
            }
        )
    out = pd.DataFrame(rows)
    out["_tier_order"] = out["Final_analysis_tier"].map(TIER_ORDER).fillna(99)
    out = out.sort_values(
        ["_tier_order", "Group_sig_p05_n", "Continuous_sig_p05_n", "Median_C_index"],
        ascending=[True, False, False, False],
    ).drop(columns=["_tier_order"])
    return out


def build_matrix(pairwise):
    ok = pairwise[pairwise["status"].eq("ok")].copy()
    rows = []
    for (sig, gse), sub in ok.groupby(["Signature_ID", "GSE"], sort=False):
        row = sub.iloc[0]
        rows.append(
            {
                "Final_analysis_tier": row["Final_analysis_tier"],
                "Article_ID": row["Article_ID"],
                "Signature_ID": sig,
                "GSE": gse,
                "N": row["N"],
                "Events": row["Events"],
                "C_index": row["C_index"],
                "Group_HR": row["Cox_high_vs_low_HR"],
                "Group_HR_CI95": ci_text(row["Cox_high_vs_low_CI95_low"], row["Cox_high_vs_low_CI95_high"]),
                "Group_p": row["Cox_high_vs_low_p"],
                "Continuous_HR_per_SD": row["Cox_continuous_HR_per_SD"],
                "Continuous_p": row["Cox_continuous_p"],
                "Cutoff_rule_used": row.get("Cutoff_rule_used", ""),
                "Analysis_source": row.get("Analysis_source", ""),
                "GSE_was_original_validation": row.get("GSE_was_original_validation", False),
            }
        )
    return pd.DataFrame(rows)


def ci_text(lo, hi):
    if pd.isna(lo) or pd.isna(hi):
        return ""
    return f"{lo:.2f}-{hi:.2f}"


def write_workbook(signature, pairwise, matrix, meta):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    main_ok = pairwise[pairwise["Is_main_analysis"] & pairwise["status"].eq("ok")]
    main_unique = main_ok.drop_duplicates("GSE")

    def add_sheet(name, df, freeze=True):
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            ws.append(list(row))
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if freeze:
            ws.freeze_panes = "A2"
        for idx, col in enumerate(df.columns, start=1):
            width = 16
            if any(token in col for token in ["Signature", "Title", "GSEs", "caveats", "Role", "CI95", "validation"]):
                width = 34
            if col in {"Final_analysis_tier", "Key_caveats_for_manuscript"}:
                width = 50
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.auto_filter.ref = ws.dimensions
        return ws

    summary_rows = pd.DataFrame(
        [
            {"Metric": "Main-analysis signatures", "Value": int(signature["Main_or_sensitivity"].eq("main").sum())},
            {"Metric": "Main-analysis GEO tests", "Value": int(main_ok.shape[0])},
            {"Metric": "Unique main-analysis GEO cohorts", "Value": int(main_unique.shape[0])},
            {"Metric": "Unique main-analysis OS-complete samples", "Value": int(main_unique["N"].sum())},
            {"Metric": "Unique main-analysis OS events", "Value": int(main_unique["Events"].sum())},
            {
                "Metric": "Main-analysis grouped Cox p<0.05 tests",
                "Value": int(main_ok["Group_sig_p05"].sum()),
            },
            {
                "Metric": "Main-analysis continuous Cox p<0.05 tests",
                "Value": int(main_ok["Continuous_sig_p05"].sum()),
            },
            {"Metric": "Strict reported-cutoff signatures", "Value": int(signature["Final_analysis_tier"].eq("Tier_1_strict_reported_cutoff").sum())},
            {"Metric": "Primary median-fallback signatures", "Value": int(signature["Final_analysis_tier"].eq("Tier_2_primary_median_fallback").sum())},
            {"Metric": "Sensitivity signatures", "Value": int(signature["Main_or_sensitivity"].eq("sensitivity").sum())},
        ]
    )
    add_sheet("Summary", summary_rows, freeze=False)
    decision = build_decision_ranking(signature)
    add_sheet("Decision_Ranking", decision)
    add_sheet("Signature_Evidence", signature)
    add_sheet("Pooled_HR", meta)
    add_sheet("Pairwise_Evidence", pairwise)
    add_sheet("Cohort_Matrix_Long", matrix)

    main_matrix = matrix[matrix["Final_analysis_tier"].isin(MAIN_TIERS)].copy()
    pivot = main_matrix.pivot_table(index="Article_ID", columns="GSE", values="Group_HR", aggfunc="first")
    pivot = pivot.reindex(columns=GSE_ORDER)
    pivot = pivot.reset_index()
    add_sheet("Main_HR_Matrix", pivot)

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    wb.save(OUT_XLSX)


def fmt_p(p):
    if pd.isna(p):
        return "NA"
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def write_results_draft(signature, pairwise):
    main = signature[signature["Main_or_sensitivity"].eq("main")].copy()
    strict = main[main["Final_analysis_tier"].eq("Tier_1_strict_reported_cutoff")]
    median = main[main["Final_analysis_tier"].eq("Tier_2_primary_median_fallback")]
    main_pair = pairwise[pairwise["Is_main_analysis"] & pairwise["status"].eq("ok")]
    main_unique = main_pair.drop_duplicates("GSE")
    lines = [
        "# Draft results text for GEO external portability",
        "",
        "This is a working manuscript draft generated from the current reproducibility and GEO validation tables.",
        "",
        "## Main findings",
        "",
        (
            f"Among the reproducible LUAD transcriptomic prognostic signatures, {len(main)} signatures were advanced "
            f"to the primary GEO portability analysis. One signature ({'; '.join(strict['Article_ID'].astype(str))}) "
            "could be assessed using a reported numeric cutoff, whereas "
            f"{len(median)} signatures required cohort-specific median dichotomisation because no portable external "
            "cutoff could be confirmed."
        ),
        "",
        (
            f"Across the primary analysis set, {len(main_pair)} signature-cohort tests were evaluable across "
            f"{len(main_unique)} GEO cohorts ({int(main_unique['N'].sum())} OS-complete LUAD-like samples and "
            f"{int(main_unique['Events'].sum())} OS events). Grouped high-vs-low Cox tests were significant in "
            f"{int(main_pair['Group_sig_p05'].sum())}/{len(main_pair)} tests, and continuous per-SD Cox tests were "
            f"significant in {int(main_pair['Continuous_sig_p05'].sum())}/{len(main_pair)} tests."
        ),
        "",
        "## Signature-level interpretation",
        "",
    ]
    for _, row in main.iterrows():
        pooled = ""
        if pd.notna(row.get("Pooled_HR_fixed")):
            pooled = (
                f" fixed-effect descriptive pooled HR {row['Pooled_HR_fixed']:.2f} "
                f"(95% CI {row['Pooled_HR_CI95_low']:.2f}-{row['Pooled_HR_CI95_high']:.2f}, "
                f"p={fmt_p(row['Pooled_HR_p'])}, I2={row['Meta_I2_percent']:.1f}%)."
            )
        lines.append(
            f"- {row['Article_ID']} ({row['Signature_ID']}): {row['GEO_tests_ok']} GEO cohorts; "
            f"grouped p<0.05 in {row['Group_sig_p05_n']}/{row['GEO_tests_ok']}; "
            f"continuous p<0.05 in {row['Continuous_sig_p05_n']}/{row['GEO_tests_ok']}; "
            f"median C-index {row['Median_C_index']:.3f}; median grouped HR {row['Median_group_HR']:.2f};"
            f"{pooled} Caveat: {row['Key_caveats_for_manuscript']}"
        )
    lines.extend(
        [
            "",
            "## Reporting caution",
            "",
            "Median-fallback analyses should be described as score-portability or directionality analyses, not as validation of the originally published risk-group cutoff. A097 is a clear example: the uploaded PDF confirms a TCGA/training cutoff of 1.66, but applying that fixed cutoff to processed GEO matrices collapses all tested cohorts to one group.",
            "",
            "GSE50081 remains flagged for final survival-time-unit verification before freezing manuscript tables.",
            "",
        ]
    )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def build_decision_ranking(signature):
    rows = []
    for _, row in signature.iterrows():
        article = row["Article_ID"]
        tier = row["Final_analysis_tier"]
        if tier == "Tier_1_strict_reported_cutoff":
            decision = "Core strict-cutoff analysis"
            rationale = "Only signature with reproducible formula and portable reported numeric cutoff in the current GEO screen."
        elif article in {"A129", "A067"}:
            decision = "Headline score-portability candidate"
            rationale = "Strongest balance of cohort count, grouped significance, continuous significance, and direction consistency under median fallback."
        elif article == "A097":
            decision = "Headline/supporting score-portability with explicit cutoff caveat"
            rationale = "Strong GEO score signal, but uploaded PDF confirms the fixed TCGA cutoff is not portable to processed GEO matrices."
        elif tier == "Tier_2_primary_median_fallback":
            decision = "Supporting primary score-portability candidate"
            rationale = "Adequate cross-GEO evidence, but should be presented after the stronger median-fallback candidates."
        elif tier == "Tier_3_secondary_median_fallback":
            decision = "Supplementary median-fallback sensitivity"
            rationale = "Useful for robustness and breadth, but not strong enough for headline claims."
        elif tier == "Tier_4_continuous_score_support":
            decision = "Supplementary continuous-score evidence"
            rationale = "Continuous score has support; grouped median-fallback evidence is not robust enough for main text."
        elif tier == "Tier_5_manual_override_sensitivity":
            decision = "Formula/symbol-correction sensitivity only"
            rationale = "Requires manual formula or gene-symbol intervention; do not mix with automated first-pass main analysis."
        else:
            decision = "Do not advance"
            rationale = "Insufficient manuscript-ready external portability evidence."
        rows.append(
            {
                "Decision": decision,
                "Article_ID": article,
                "Signature_ID": row["Signature_ID"],
                "Final_analysis_tier": tier,
                "GEO_tests_ok": row["GEO_tests_ok"],
                "Group_sig_p05_n": row["Group_sig_p05_n"],
                "Continuous_sig_p05_n": row["Continuous_sig_p05_n"],
                "Median_C_index": row["Median_C_index"],
                "Median_group_HR": row["Median_group_HR"],
                "Pooled_HR_fixed": row.get("Pooled_HR_fixed", np.nan),
                "Pooled_HR_CI95_low": row.get("Pooled_HR_CI95_low", np.nan),
                "Pooled_HR_CI95_high": row.get("Pooled_HR_CI95_high", np.nan),
                "Meta_I2_percent": row.get("Meta_I2_percent", np.nan),
                "Rationale": rationale,
                "Manuscript_caveat": row["Key_caveats_for_manuscript"],
            }
        )
    out = pd.DataFrame(rows)
    decision_order = {
        "Core strict-cutoff analysis": 1,
        "Headline score-portability candidate": 2,
        "Headline/supporting score-portability with explicit cutoff caveat": 3,
        "Supporting primary score-portability candidate": 4,
        "Supplementary median-fallback sensitivity": 5,
        "Supplementary continuous-score evidence": 6,
        "Formula/symbol-correction sensitivity only": 7,
        "Do not advance": 8,
    }
    out["_order"] = out["Decision"].map(decision_order).fillna(99)
    out = out.sort_values(
        ["_order", "Group_sig_p05_n", "Continuous_sig_p05_n", "Median_C_index"],
        ascending=[True, False, False, False],
    ).drop(columns=["_order"])
    return out


def svg_escape(text):
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def color_for_hr(hr, p):
    if pd.isna(hr):
        return "#F3F4F6"
    if hr >= 1:
        intensity = min(1.0, (math.log(hr) / math.log(3.0)))
        base = np.array([178, 24, 43])
        light = np.array([253, 219, 199])
    else:
        intensity = min(1.0, abs(math.log(hr)) / math.log(3.0))
        base = np.array([33, 102, 172])
        light = np.array([209, 229, 240])
    rgb = (light * (1 - intensity) + base * intensity).astype(int)
    if pd.isna(p) or p >= 0.05:
        rgb = (rgb * 0.55 + np.array([245, 245, 245]) * 0.45).astype(int)
    return "#" + "".join(f"{v:02X}" for v in rgb)


def write_heatmap(matrix):
    main = matrix[matrix["Final_analysis_tier"].isin(MAIN_TIERS)].copy()
    sigs = list(dict.fromkeys(main["Article_ID"]))
    cell_w, cell_h = 82, 34
    left, top = 125, 58
    width = left + cell_w * len(GSE_ORDER) + 35
    height = top + cell_h * len(sigs) + 70
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<text x="18" y="24" font-family="Arial" font-size="16" font-weight="700">Main GEO portability: grouped HR by cohort</text>',
        '<text x="18" y="44" font-family="Arial" font-size="11" fill="#4B5563">Color shows high-vs-low HR; darker cells have larger effect, faded cells are p>=0.05.</text>',
    ]
    for j, gse in enumerate(GSE_ORDER):
        x = left + j * cell_w + cell_w / 2
        parts.append(f'<text x="{x}" y="{top - 12}" text-anchor="middle" font-family="Arial" font-size="10">{gse}</text>')
    lookup = {(r.Article_ID, r.GSE): r for r in main.itertuples(index=False)}
    for i, sig in enumerate(sigs):
        y = top + i * cell_h
        parts.append(f'<text x="{left - 10}" y="{y + 22}" text-anchor="end" font-family="Arial" font-size="11" font-weight="700">{svg_escape(sig)}</text>')
        for j, gse in enumerate(GSE_ORDER):
            x = left + j * cell_w
            row = lookup.get((sig, gse))
            hr = getattr(row, "Group_HR", np.nan) if row is not None else np.nan
            p = getattr(row, "Group_p", np.nan) if row is not None else np.nan
            fill = color_for_hr(hr, p)
            label = "" if pd.isna(hr) else f"{hr:.2f}" + ("*" if pd.notna(p) and p < 0.05 else "")
            parts.append(f'<rect x="{x}" y="{y}" width="{cell_w - 2}" height="{cell_h - 2}" fill="{fill}" stroke="#FFFFFF"/>')
            parts.append(f'<text x="{x + cell_w/2}" y="{y + 21}" text-anchor="middle" font-family="Arial" font-size="10">{label}</text>')
    parts.append(f'<text x="{left}" y="{height - 26}" font-family="Arial" font-size="10" fill="#4B5563">* Cox high-vs-low p&lt;0.05. Empty cells indicate non-evaluable signature-cohort tests.</text>')
    parts.append("</svg>")
    OUT_HEATMAP.write_text("\n".join(parts), encoding="utf-8")


def write_forest(signature):
    main = signature[signature["Main_or_sensitivity"].eq("main") & signature["Pooled_HR_fixed"].notna()].copy()
    main = main.sort_values(["Final_analysis_tier", "Pooled_HR_fixed"], ascending=[True, False])
    left, top, row_h = 170, 58, 34
    plot_w = 430
    width = left + plot_w + 230
    height = top + row_h * len(main) + 72
    xmin, xmax = 0.6, 4.0
    log_min, log_max = math.log(xmin), math.log(xmax)

    def xmap(v):
        return left + (math.log(max(xmin, min(xmax, v))) - log_min) / (log_max - log_min) * plot_w

    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<text x="18" y="24" font-family="Arial" font-size="16" font-weight="700">Descriptive fixed-effect pooled HR across GEO cohorts</text>',
        '<text x="18" y="44" font-family="Arial" font-size="11" fill="#4B5563">Pooled across evaluable grouped high-vs-low Cox tests; interpret as a portability summary, not a definitive meta-analysis.</text>',
    ]
    axis_y = top + row_h * len(main) + 14
    for tick in [0.75, 1.0, 1.5, 2.0, 3.0, 4.0]:
        x = xmap(tick)
        parts.append(f'<line x1="{x}" y1="{top - 8}" x2="{x}" y2="{axis_y}" stroke="#E5E7EB"/>')
        parts.append(f'<text x="{x}" y="{axis_y + 15}" text-anchor="middle" font-family="Arial" font-size="10">{tick:g}</text>')
    parts.append(f'<line x1="{xmap(1)}" y1="{top - 12}" x2="{xmap(1)}" y2="{axis_y}" stroke="#111827" stroke-dasharray="3 3"/>')
    for i, row in enumerate(main.itertuples(index=False)):
        y = top + i * row_h
        mid_y = y + 17
        lo, hr, hi = row.Pooled_HR_CI95_low, row.Pooled_HR_fixed, row.Pooled_HR_CI95_high
        parts.append(f'<text x="{left - 12}" y="{mid_y + 4}" text-anchor="end" font-family="Arial" font-size="11" font-weight="700">{svg_escape(row.Article_ID)}</text>')
        parts.append(f'<line x1="{xmap(lo)}" y1="{mid_y}" x2="{xmap(hi)}" y2="{mid_y}" stroke="#374151" stroke-width="2"/>')
        parts.append(f'<circle cx="{xmap(hr)}" cy="{mid_y}" r="5" fill="#B2182B"/>')
        p = fmt_p(row.Pooled_HR_p)
        label = f'{hr:.2f} ({lo:.2f}-{hi:.2f}), p={p}, I2={row.Meta_I2_percent:.0f}%'
        parts.append(f'<text x="{left + plot_w + 18}" y="{mid_y + 4}" font-family="Arial" font-size="10">{svg_escape(label)}</text>')
    parts.append(f'<text x="{left + plot_w / 2}" y="{height - 12}" text-anchor="middle" font-family="Arial" font-size="10">Hazard ratio, log scale</text>')
    parts.append("</svg>")
    OUT_FOREST.write_text("\n".join(parts), encoding="utf-8")


def main():
    selection = pd.read_csv(SELECTION)
    selection = selection[selection["Final_analysis_tier"].isin(set(TIER_ORDER))].copy()
    selection["_tier_order"] = selection["Final_analysis_tier"].map(TIER_ORDER)
    selection = selection.sort_values(["_tier_order", "Group_sig_p05_n", "Continuous_sig_p05_n"], ascending=[True, False, False]).drop(columns=["_tier_order"])

    pairwise = load_pairwise(selection)
    signature = build_signature_evidence(selection, pairwise)
    matrix = build_matrix(pairwise)
    meta = signature[
        [
            "Final_analysis_tier",
            "Article_ID",
            "Signature_ID",
            "GEO_tests_ok",
            "Pooled_HR_fixed",
            "Pooled_HR_CI95_low",
            "Pooled_HR_CI95_high",
            "Pooled_HR_p",
            "Meta_k",
            "Meta_I2_percent",
            "Cutoff_basis_for_GEO",
        ]
    ].copy()
    decision = build_decision_ranking(signature)

    signature.to_csv(OUT_SIGNATURE, index=False, encoding="utf-8")
    pairwise.to_csv(OUT_PAIRWISE, index=False, encoding="utf-8")
    matrix.to_csv(OUT_MATRIX, index=False, encoding="utf-8")
    meta.to_csv(OUT_META, index=False, encoding="utf-8")
    decision.to_csv(OUT_DECISION, index=False, encoding="utf-8")
    write_workbook(signature, pairwise, matrix, meta)
    write_results_draft(signature, pairwise)
    write_heatmap(matrix)
    write_forest(signature)

    main_pair = pairwise[pairwise["Is_main_analysis"] & pairwise["status"].eq("ok")]
    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_MD}")
    print(f"Wrote {OUT_HEATMAP}")
    print(f"Wrote {OUT_FOREST}")
    print(f"Main signatures: {signature['Main_or_sensitivity'].eq('main').sum()}")
    print(f"Main evaluable GEO tests: {len(main_pair)}")
    print(f"Main grouped p<0.05: {int(main_pair['Group_sig_p05'].sum())}/{len(main_pair)}")
    print(f"Main continuous p<0.05: {int(main_pair['Continuous_sig_p05'].sum())}/{len(main_pair)}")


if __name__ == "__main__":
    main()
